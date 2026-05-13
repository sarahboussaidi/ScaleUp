import json
import re
from dataclasses import dataclass
from pathlib import Path
from typing import Dict, List, Optional, Tuple

from docx import Document


@dataclass
class FieldPrediction:
    value: str
    confidence: float
    source: str


class IntelligentNDAFiller:
    """Infer NDA placeholders from unstructured text plus optional hints."""

    TITLE_PATTERN = re.compile(
        r"\b(CEO|CTO|CFO|COO|Managing Director|Director|Founder|Co-Founder|Attorney)\b",
        re.IGNORECASE,
    )
    EMAIL_PATTERN = re.compile(r"[A-Za-z0-9._%+-]+@[A-Za-z0-9.-]+\.[A-Za-z]{2,}")
    DATE_PATTERN = re.compile(
        r"\b(\d{1,2}\s+(?:January|February|March|April|May|June|July|August|September|October|November|December)\s+\d{4})\b",
        re.IGNORECASE,
    )
    REGISTRY_PATTERN = re.compile(r"\b\d{6,12}\b")
    TAX_PATTERN = re.compile(r"\b[A-Z]{2}\d{6,12}\b")
    ORG_PATTERN = re.compile(
        r"\b([A-Z][A-Za-z0-9&.,' -]{2,}?(?:\s+(?:O\u00dc|OU|Ltd|LLC|Inc|GmbH|SAS|PLC|B\.V\.|BV)))\b"
    )
    PERSON_PATTERN = re.compile(r"\b([A-Z][a-z]+\s+[A-Z][a-z]+)\b")

    def __init__(self, template_path: str):
        self.template_path = Path(template_path)

    def _unique(self, values: List[str]) -> List[str]:
        seen = set()
        out = []
        for item in values:
            key = item.strip()
            if not key:
                continue
            low = key.lower()
            if low in seen:
                continue
            seen.add(low)
            out.append(key)
        return out

    def _extract_orgs(self, text: str) -> List[str]:
        return self._unique(self.ORG_PATTERN.findall(text))

    def _extract_people(self, text: str) -> List[str]:
        candidates = self._unique(self.PERSON_PATTERN.findall(text))
        return [name for name in candidates if name.lower() not in {"month yyyy"}]

    def _extract_titles(self, text: str) -> List[str]:
        return self._unique([match.group(0) for match in self.TITLE_PATTERN.finditer(text)])

    def _extract_date(self, text: str) -> Optional[str]:
        match = self.DATE_PATTERN.search(text)
        return match.group(1) if match else None

    def _extract_purpose(self, text: str) -> Optional[str]:
        purpose_patterns = [
            re.compile(r"purpose\s*[:\-]\s*([^\n\.]+)", re.IGNORECASE),
            re.compile(r"regarding\s+([^\n\.]+)", re.IGNORECASE),
            re.compile(r"to\s+(discuss[^\n\.]+)", re.IGNORECASE),
        ]
        for pattern in purpose_patterns:
            found = pattern.search(text)
            if found:
                return found.group(1).strip().rstrip(".")
        return None

    def _extract_labeled_pair(self, text: str, labels: Tuple[str, str], pattern: re.Pattern) -> List[str]:
        values: List[str] = []
        for label in labels:
            regex = re.compile(label + r"\s*[:\-]\s*([^\n,;]+)", re.IGNORECASE)
            for match in regex.finditer(text):
                chunk = match.group(1).strip()
                values.extend(pattern.findall(chunk))
        if not values:
            values = pattern.findall(text)
        return self._unique(values)

    def _extract_registry(self, text: str) -> List[str]:
        return self._extract_labeled_pair(
            text,
            (r"party\s*1\s*registry", r"party\s*2\s*registry"),
            self.REGISTRY_PATTERN,
        )

    def _extract_tax_ids(self, text: str) -> List[str]:
        return self._extract_labeled_pair(
            text,
            (r"party\s*1\s*tax", r"party\s*2\s*tax"),
            self.TAX_PATTERN,
        )

    def _extract_addresses(self, text: str) -> List[str]:
        explicit = []
        for label in ("party 1 address", "party 2 address", "address"):
            regex = re.compile(label + r"\s*[:\-]\s*([^\n]+)", re.IGNORECASE)
            explicit.extend([m.group(1).strip() for m in regex.finditer(text)])
        explicit = self._unique(explicit)
        if explicit:
            return explicit

        fallback = []
        for line in text.splitlines():
            line = line.strip()
            if "," in line and len(line) < 120 and any(c.isalpha() for c in line):
                fallback.append(line)
        return self._unique(fallback)

    def _extract_emails(self, text: str) -> List[str]:
        return self._unique(self.EMAIL_PATTERN.findall(text))

    def _placeholder_order_for_insert(self) -> List[str]:
        canonical = [
            "registry_1",
            "registry_2",
            "address_1",
            "address_2",
            "email_1",
            "email_2",
            "tax_1",
            "tax_2",
        ]

        if not self.template_path.exists():
            return canonical

        doc = Document(str(self.template_path))
        contexts = []

        def _scan_line(line: str) -> None:
            low = line.lower()
            if "[insert]" not in low:
                return
            count = low.count("[insert]")
            if "registry" in low or "commercial register" in low:
                contexts.extend(["registry"] * count)
            elif "address" in low or "registered office" in low or "domicile" in low:
                contexts.extend(["address"] * count)
            elif "e-mail" in low or "email" in low:
                contexts.extend(["email"] * count)
            elif "vat" in low or "tax" in low:
                contexts.extend(["tax"] * count)
            else:
                contexts.extend(["unknown"] * count)

        for paragraph in doc.paragraphs:
            _scan_line(paragraph.text)

        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    for paragraph in cell.paragraphs:
                        _scan_line(paragraph.text)

        counters = {"registry": 0, "address": 0, "email": 0, "tax": 0, "unknown": 0}
        ordered = []
        for kind in contexts:
            counters[kind] += 1
            ordered.append(f"{kind}_{counters[kind]}")

        # If context is clearly skewed (common in single-line legal boilerplate),
        # fall back to canonical legal-field order for deterministic quality.
        classified = {k: v for k, v in counters.items() if k != "unknown"}
        max_classified = max(classified.values(), default=0)
        if not ordered or (len(ordered) >= 8 and max_classified > len(ordered) * 0.5):
            return canonical

        return ordered

    def _build_insert_values(
        self,
        registry: List[str],
        addresses: List[str],
        emails: List[str],
        tax_ids: List[str],
    ) -> Tuple[List[str], List[str]]:
        order = self._placeholder_order_for_insert()

        pools = {
            "registry": registry,
            "address": addresses,
            "email": emails,
            "tax": tax_ids,
            "unknown": registry + addresses + emails + tax_ids,
        }
        ptr = {"registry": 0, "address": 0, "email": 0, "tax": 0, "unknown": 0}

        values = []
        warnings = []

        for key in order:
            kind = key.split("_", 1)[0]
            if kind not in pools:
                kind = "unknown"

            if ptr[kind] >= len(pools[kind]):
                values.append("MISSING")
                warnings.append(f"Could not infer value for [insert] slot '{key}'.")
                continue

            values.append(pools[kind][ptr[kind]])
            ptr[kind] += 1

        return values, warnings

    def infer(self, brief_text: str, hints: Optional[Dict[str, str]] = None) -> Dict[str, object]:
        hints = hints or {}

        orgs = self._extract_orgs(brief_text)
        people = self._extract_people(brief_text)
        titles = self._extract_titles(brief_text)
        emails = self._extract_emails(brief_text)
        registries = self._extract_registry(brief_text)
        addresses = self._extract_addresses(brief_text)
        tax_ids = self._extract_tax_ids(brief_text)
        date_value = self._extract_date(brief_text)
        purpose_value = self._extract_purpose(brief_text)

        p1 = hints.get("party1_name") or (orgs[0] if len(orgs) >= 1 else "Party One")
        p2 = hints.get("party2_name") or (orgs[1] if len(orgs) >= 2 else "Party Two")

        rep1 = hints.get("rep1_name") or (people[0] if len(people) >= 1 else "Representative One")
        rep2 = hints.get("rep2_name") or (people[1] if len(people) >= 2 else "Representative Two")

        title1 = hints.get("rep1_title") or (titles[0] if len(titles) >= 1 else "Director")
        title2 = hints.get("rep2_title") or (titles[1] if len(titles) >= 2 else "Director")

        date_final = hints.get("date") or date_value or "31 March 2026"
        purpose_final = hints.get("purpose") or purpose_value or "evaluating a potential business collaboration"

        other_party_type = hints.get("other_party_domicile_type")
        if not other_party_type:
            addr = addresses[1] if len(addresses) > 1 else (addresses[0] if addresses else "its registered location")
            other_party_type = f"a company with registered office in {addr}"

        insert_values, insert_warnings = self._build_insert_values(registries, addresses, emails, tax_ids)

        data = {
            "[dd Month YYYY]": date_final,
            "[Name of 1st Party]": p1,
            "[Name of 2nd Party]": p2,
            "[Name of one Party]": p1,
            "[Name of the other Party]": p2,
            "[describe the purpose]": purpose_final,
            "[other Party's domicile and type]": other_party_type,
            "[Representative's name]": [rep1, rep2],
            "[Representative's title]": [title1, title2],
            "[insert]": insert_values,
        }

        predictions = {
            "party1_name": FieldPrediction(p1, 0.92 if len(orgs) >= 1 else 0.55, "ORG extraction + hints"),
            "party2_name": FieldPrediction(p2, 0.90 if len(orgs) >= 2 else 0.55, "ORG extraction + hints"),
            "representative_1": FieldPrediction(rep1, 0.85 if len(people) >= 1 else 0.45, "PERSON extraction + hints"),
            "representative_2": FieldPrediction(rep2, 0.82 if len(people) >= 2 else 0.45, "PERSON extraction + hints"),
            "purpose": FieldPrediction(purpose_final, 0.80 if purpose_value else 0.50, "Pattern extraction + fallback"),
            "date": FieldPrediction(date_final, 0.95 if date_value else 0.60, "Date extraction + fallback"),
        }

        warnings = []
        if "MISSING" in insert_values:
            warnings.append("Some [insert] placeholders could not be inferred and were set to 'MISSING'.")
        warnings.extend(insert_warnings)

        report = {
            "predictions": {
                key: {
                    "value": pred.value,
                    "confidence": pred.confidence,
                    "source": pred.source,
                }
                for key, pred in predictions.items()
            },
            "extracted_candidates": {
                "organizations": orgs,
                "people": people,
                "titles": titles,
                "emails": emails,
                "registries": registries,
                "addresses": addresses,
                "tax_ids": tax_ids,
            },
            "warnings": warnings,
            "placeholder_data": data,
        }

        return report

    def save_report(self, report: Dict[str, object], output_path: str) -> None:
        out_file = Path(output_path)
        out_file.parent.mkdir(parents=True, exist_ok=True)
        out_file.write_text(json.dumps(report, indent=2), encoding="utf-8")
