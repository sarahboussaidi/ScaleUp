import re
from pathlib import Path

from docx import Document

class TemplateDocumentGenerator:
    def __init__(self):
        pass

    def _remove_paragraph(self, paragraph):
        parent = paragraph._element.getparent()
        if parent is not None:
            parent.remove(paragraph._element)

    def _paragraph_matches_clause(self, paragraph, clause_numbers):
        text = (paragraph.text or "").strip()
        match = re.match(r"^(\d+)\.\s+", text)
        if not match:
            return False
        return int(match.group(1)) in clause_numbers

    def _remove_clause_blocks(self, doc, clause_numbers):
        if not clause_numbers:
            return

        target_numbers = {int(number) for number in clause_numbers if str(number).isdigit()}
        if not target_numbers:
            return

        paragraphs = list(doc.paragraphs)
        ranges = []
        for idx, paragraph in enumerate(paragraphs):
            text = (paragraph.text or "").strip()
            match = re.match(r"^(\d+)\.\s+", text)
            if not match:
                continue
            clause_number = int(match.group(1))
            if clause_number not in target_numbers:
                continue

            end_idx = len(paragraphs)
            for next_idx in range(idx + 1, len(paragraphs)):
                next_text = (paragraphs[next_idx].text or "").strip()
                if re.match(r"^\d+\.\s+", next_text) or re.match(r"^signatures?:?\s*$", next_text, re.IGNORECASE):
                    end_idx = next_idx
                    break
            ranges.append((idx, end_idx))

        for start_idx, end_idx in reversed(ranges):
            for paragraph in paragraphs[start_idx:end_idx]:
                self._remove_paragraph(paragraph)

    def _replace_token_in_paragraph(self, paragraph, token, replacement):
        runs = list(paragraph.runs)
        if not runs:
            return False

        full_text = "".join(run.text for run in runs)
        start_index = full_text.find(token)
        if start_index < 0:
            return False

        end_index = start_index + len(token)
        cursor = 0
        start_run_index = None
        start_run_offset = 0
        end_run_index = None
        end_run_offset = 0

        for idx, run in enumerate(runs):
            run_text = run.text
            run_end = cursor + len(run_text)

            if start_run_index is None and start_index < run_end:
                start_run_index = idx
                start_run_offset = start_index - cursor

            if end_index <= run_end:
                end_run_index = idx
                end_run_offset = end_index - cursor
                break

            cursor = run_end

        if start_run_index is None or end_run_index is None:
            return False

        if start_run_index == end_run_index:
            run = runs[start_run_index]
            run.text = run.text[:start_run_offset] + replacement + run.text[end_run_offset:]
            return True

        start_run = runs[start_run_index]
        end_run = runs[end_run_index]
        start_prefix = start_run.text[:start_run_offset]
        end_suffix = end_run.text[end_run_offset:]
        start_run.text = start_prefix + replacement

        for idx in range(start_run_index + 1, end_run_index):
            runs[idx].text = ""

        end_run.text = end_suffix
        return True

    def _replace_token_everywhere(self, paragraph, token, replacement):
        replaced = False
        while token in paragraph.text:
            if not self._replace_token_in_paragraph(paragraph, token, replacement):
                break
            replaced = True
        return replaced

    # Replace text inside paragraphs
    def replace_text(self, paragraph, data, sequence_state):
        # Replace longer tokens first so wrappers like
        # [OPTION 2: ... [insert] ...] are processed before inner [insert] tokens.
        ordered_keys = sorted(data.keys(), key=len, reverse=True)

        for key in ordered_keys:
            value = data[key]
            if key not in paragraph.text:
                continue

            if isinstance(value, (list, tuple)):
                # Replace repeated placeholders one-by-one in document order.
                while key in paragraph.text and sequence_state[key] < len(value):
                    replacement = str(value[sequence_state[key]])
                    if not self._replace_token_in_paragraph(paragraph, key, replacement):
                        break
                    sequence_state[key] += 1
            else:
                self._replace_token_everywhere(paragraph, key, str(value))

    # Replace everywhere in document
    def replace_all(self, doc, data, sequence_state):
        # Paragraphs
        for paragraph in doc.paragraphs:
            self.replace_text(paragraph, data, sequence_state)

        # Tables
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    for paragraph in cell.paragraphs:
                        self.replace_text(paragraph, data, sequence_state)

    def _collect_document_text(self, doc):
        parts = []

        for paragraph in doc.paragraphs:
            parts.append(paragraph.text)

        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    for paragraph in cell.paragraphs:
                        parts.append(paragraph.text)

        return "\n".join(parts)

    def _find_bracket_placeholders(self, doc):
        text = self._collect_document_text(doc)
        return sorted(set(re.findall(r"\[[^\]]+\]", text)))

    def _strip_all_square_brackets(self, doc):
        for paragraph in doc.paragraphs:
            for run in paragraph.runs:
                if "[" in run.text or "]" in run.text:
                    run.text = run.text.replace("[", "").replace("]", "")

        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    for paragraph in cell.paragraphs:
                        for run in paragraph.runs:
                            if "[" in run.text or "]" in run.text:
                                run.text = run.text.replace("[", "").replace("]", "")

    # Generate final document
    def generate(self, template_path, output_path, data, strict=True, strip_bracket_artifacts=True, remove_clause_numbers=None):
        template_file = Path(template_path)
        output_file = Path(output_path)

        if not template_file.exists():
            raise FileNotFoundError(f"Template not found: {template_path}")

        doc = Document(template_path)

        if remove_clause_numbers:
            self._remove_clause_blocks(doc, remove_clause_numbers)

        sequence_state = {
            key: 0 for key, value in data.items() if isinstance(value, (list, tuple))
        }

        self.replace_all(doc, data, sequence_state)

        unresolved = self._find_bracket_placeholders(doc)
        unused_sequence_values = [
            key
            for key, idx in sequence_state.items()
            if idx < len(data[key])
        ]

        if strict and (unresolved or unused_sequence_values):
            details = []

            if unresolved:
                details.append(
                    "Unreplaced placeholders: " + ", ".join(unresolved)
                )

            if unused_sequence_values:
                details.append(
                    "Unused sequence values for placeholders: "
                    + ", ".join(unused_sequence_values)
                )

            raise ValueError("Template replacement incomplete. " + " | ".join(details))

        # Safety pass: remove any lingering square-bracket artifacts from the final text.
        # This keeps outputs user-friendly when templates contain non-standard bracket remnants.
        if strip_bracket_artifacts:
            full_text = self._collect_document_text(doc)
            if "[" in full_text or "]" in full_text:
                self._strip_all_square_brackets(doc)

        output_file.parent.mkdir(parents=True, exist_ok=True)
        doc.save(output_path)

        return output_path