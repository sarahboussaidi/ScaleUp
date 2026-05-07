import json
import os
import re
from pathlib import Path
from typing import Any, Dict, List, Optional

import httpx
from docx import Document
try:
    from dotenv import load_dotenv
except Exception:
    def load_dotenv(*_args: Any, **_kwargs: Any) -> bool:
        return False

from objective_registry import get_selected_value

try:
    import torch
except Exception:
    torch = None

try:
    from transformers import AutoModelForCausalLM, AutoTokenizer
except Exception:
    AutoModelForCausalLM = None
    AutoTokenizer = None


load_dotenv(Path(__file__).resolve().parent / ".env")


class LegalDocumentIntelligence:
    LEGAL_RULES = [
        {
            "category": "confidentiality",
            "title": "Confidentiality / NDA",
            "keywords": ["confidential", "non-disclosure", "nda", "secret", "proprietary", "confidential information"],
        },
        {
            "category": "purpose",
            "title": "Purpose / Scope",
            "keywords": ["purpose", "scope", "use solely", "solely for", "evaluation", "discussion"],
        },
        {
            "category": "term",
            "title": "Term / Duration",
            "keywords": ["term", "duration", "effective date", "expires", "commence", "start date"],
        },
        {
            "category": "termination",
            "title": "Termination",
            "keywords": ["terminate", "termination", "notice period", "without notice", "expire", "end of"],
        },
        {
            "category": "governing_law",
            "title": "Governing law / Jurisdiction",
            "keywords": ["governing law", "jurisdiction", "venue", "courts", "dispute", "arbitration"],
        },
        {
            "category": "liability",
            "title": "Liability / Indemnity",
            "keywords": ["liability", "indemnity", "indemnification", "damages", "limitation of liability", "losses"],
        },
        {
            "category": "ip",
            "title": "Intellectual property",
            "keywords": ["intellectual property", "ip", "assignment", "ownership", "license", "source code", "invention"],
        },
        {
            "category": "payment",
            "title": "Payment / Fees",
            "keywords": ["payment", "fee", "fees", "invoice", "consideration", "price", "remuneration"],
        },
        {
            "category": "data_protection",
            "title": "Data protection / Privacy",
            "keywords": ["data protection", "privacy", "personal data", "gdpr", "processor", "controller"],
        },
        {
            "category": "assignment",
            "title": "Assignment / Transfer",
            "keywords": ["assignment", "assign", "transfer", "subcontract", "novation"],
        },
        {
            "category": "non_compete",
            "title": "Non-compete / Non-solicit",
            "keywords": ["non-compete", "non compete", "non-solicit", "non solicit", "restrict", "solicit"],
        },
    ]

    STARTUP_RULES = [
        {
            "category": "fundraising",
            "title": "Fundraising / Investment",
            "keywords": ["seed", "series a", "series b", "fundraising", "investment", "investor", "convertible note", "safe", "valuation"],
        },
        {
            "category": "equity",
            "title": "Equity / Ownership",
            "keywords": ["equity", "shares", "shareholder", "cap table", "dilution", "option pool", "founder", "vesting"],
        },
        {
            "category": "governance",
            "title": "Governance / Control",
            "keywords": ["board", "voting", "consent", "approval", "observer", "control", "reserved matters"],
        },
        {
            "category": "ip_strategy",
            "title": "IP / Technology",
            "keywords": ["intellectual property", "source code", "product", "technology", "license", "open source", "assignment"],
        },
        {
            "category": "commercial",
            "title": "Commercial Terms",
            "keywords": ["subscription", "pricing", "commercial", "order form", "renewal", "service levels", "sla"],
        },
        {
            "category": "employment",
            "title": "Employment / Contractors",
            "keywords": ["employee", "contractor", "employment", "offer", "salary", "non-compete", "non-solicit"],
        },
        {
            "category": "customer_risk",
            "title": "Customer / Delivery Risk",
            "keywords": ["delivery", "milestone", "acceptance", "warranty", "support", "refund", "liability", "indemnity"],
        },
        {
            "category": "fund_use",
            "title": "Use of Proceeds / Milestones",
            "keywords": ["use of proceeds", "milestone", "milestones", "growth", "burn", "runway", "budget"],
        },
    ]

    MISSING_REVIEW_ITEMS = [
        ("confidentiality", "Confirm confidentiality scope, exceptions, and survival period."),
        ("termination", "Confirm termination trigger, notice period, and post-termination duties."),
        ("governing_law", "Confirm governing law, venue, and dispute resolution forum."),
        ("liability", "Confirm whether liability is capped, excluded, or indemnified."),
        ("ip", "Confirm intellectual property ownership, assignment, and license scope."),
        ("assignment", "Confirm transfer and assignment restrictions."),
        ("data_protection", "Confirm privacy and data protection obligations."),
        ("equity", "Confirm equity, vesting, dilution, and control mechanics if this is a startup agreement."),
    ]

    IMAGE_EXTENSIONS = {".png", ".jpg", ".jpeg", ".webp", ".bmp", ".tif", ".tiff"}
    _LOCAL_TOKENIZER = None
    _LOCAL_MODEL = None
    _LOCAL_MODEL_PATH = None

    def __init__(self, mode: Optional[str] = None) -> None:
        selected_mode = (
            mode
            or os.getenv("DOCUMENT_INTEL_MODE", "")
            or str(get_selected_value("summarization", "strategy", ""))
        ).strip().lower()
        self.mode = selected_mode or "auto"

        selected_provider = str(get_selected_value("summarization", "provider", "") or "").strip().lower()
        raw_provider = (
            selected_provider
            or os.getenv("DOCUMENT_INTEL_AI_PROVIDER", "").strip().lower()
        )
        if self.mode not in {"rule_based", "rules", "rule-only"} and raw_provider in {"off", "false", "0", "none"}:
            raise RuntimeError("Document intelligence now requires an LLM backend. Set DOCUMENT_INTEL_AI_PROVIDER to openai, ollama, or local, or use rule_based mode.")

        self.ai_provider = raw_provider or ("off" if self.mode in {"rule_based", "rules", "rule-only"} else "auto")
        self.openai_configured = any(
            os.getenv(name)
            for name in (
                "OPENAI_API_KEY",
                "OPENAI_BASE_URL",
                "DOCUMENT_INTEL_OPENAI_API_KEY",
                "DOCUMENT_INTEL_OPENAI_BASE_URL",
                "DOCUMENT_INTEL_AI_API_KEY",
                "DOCUMENT_INTEL_AI_BASE_URL",
                "OPENAI_MODEL",
                "DOCUMENT_INTEL_OPENAI_MODEL",
                "DOCUMENT_INTEL_AI_MODEL",
            )
        )
        self.ollama_configured = any(
            os.getenv(name)
            for name in (
                "OLLAMA_BASE_URL",
                "DOCUMENT_INTEL_OLLAMA_BASE_URL",
                "OLLAMA_API_KEY",
                "DOCUMENT_INTEL_OLLAMA_API_KEY",
                "OLLAMA_MODEL",
                "DOCUMENT_INTEL_OLLAMA_MODEL",
            )
        )
        self.openai_api_key = (
            os.getenv("OPENAI_API_KEY")
            or os.getenv("DOCUMENT_INTEL_OPENAI_API_KEY")
            or os.getenv("DOCUMENT_INTEL_AI_API_KEY")
            or ""
        ).strip()
        self.ollama_api_key = (
            os.getenv("OLLAMA_API_KEY")
            or os.getenv("DOCUMENT_INTEL_OLLAMA_API_KEY")
            or ""
        ).strip()
        self.openai_model = (
            os.getenv("OPENAI_MODEL")
            or os.getenv("DOCUMENT_INTEL_OPENAI_MODEL")
            or os.getenv("DOCUMENT_INTEL_AI_MODEL")
            or "gpt-4o-mini"
        ).strip()
        self.ollama_model = (
            os.getenv("OLLAMA_MODEL")
            or os.getenv("DOCUMENT_INTEL_OLLAMA_MODEL")
            or "llama3.1"
        ).strip()
        self.openai_base_url = (
            os.getenv("OPENAI_BASE_URL")
            or os.getenv("DOCUMENT_INTEL_OPENAI_BASE_URL")
            or os.getenv("DOCUMENT_INTEL_AI_BASE_URL")
            or "https://api.openai.com/v1"
        ).strip().rstrip("/")
        self.ollama_base_url = (
            os.getenv("OLLAMA_BASE_URL")
            or os.getenv("DOCUMENT_INTEL_OLLAMA_BASE_URL")
            or "http://localhost:11434/v1"
        ).strip().rstrip("/")
        selected_local_model_path = str(get_selected_value("summarization", "model_path", "") or "").strip()
        self.local_model_path = (
            os.getenv("DOCUMENT_INTEL_LOCAL_MODEL_PATH")
            or os.getenv("LOCAL_LLAMA_MODEL_PATH")
            or selected_local_model_path
            or ""
        ).strip()
        local_path = Path(self.local_model_path).expanduser() if self.local_model_path else None
        self.local_configured = bool(local_path and local_path.exists())
        self.local_max_new_tokens = int(os.getenv("DOCUMENT_INTEL_LOCAL_MAX_NEW_TOKENS", "700"))
        self.ai_model = self.openai_model if self.ai_provider in {"openai", "auto"} else self.ollama_model
        if self.ai_provider in {"local", "finetuned", "llama_local"} and self.local_model_path:
            self.ai_model = self.local_model_path
        self.ai_timeout = float(os.getenv("DOCUMENT_INTEL_AI_TIMEOUT", "45"))
        self._spacy = None
        self._ocr_analyzer = None

    def _normalize_text(self, text: str) -> str:
        cleaned = re.sub(r"\r\n?", "\n", text or "")
        cleaned = re.sub(r"[ \t]+", " ", cleaned)
        cleaned = re.sub(r"\n{3,}", "\n\n", cleaned)
        return cleaned.strip()

    def _load_docx_text(self, file_path: str) -> str:
        document = Document(file_path)
        parts: List[str] = []
        for paragraph in document.paragraphs:
            if paragraph.text.strip():
                parts.append(paragraph.text.strip())
        for table in document.tables:
            for row in table.rows:
                for cell in row.cells:
                    for paragraph in cell.paragraphs:
                        if paragraph.text.strip():
                            parts.append(paragraph.text.strip())
        return self._normalize_text("\n".join(parts))

    def _load_spacy(self):
        if self._spacy is None:
            import spacy  # type: ignore

            self._spacy = spacy.load("en_core_web_sm")
        return self._spacy

    def _load_image_text(self, file_path: str) -> str:
        if self._ocr_analyzer is None:
            from ocr.document_analyzer import DocumentAnalyzer

            self._ocr_analyzer = DocumentAnalyzer()

        result = self._ocr_analyzer.analyze(file_path)
        return self._normalize_text(str(result.get("text", "")))

    def _split_units(self, text: str) -> List[str]:
        lines = [line.strip() for line in self._normalize_text(text).splitlines()]
        units: List[str] = []
        buffer: List[str] = []

        def append_sentence_units(value: str) -> None:
            # Break very long prose lines into sentence-like units so clause
            # extraction can detect multiple obligations in one paragraph.
            candidates = [seg.strip() for seg in re.split(r"(?<=[.!?])\s+", value) if seg.strip()]
            if len(candidates) <= 1:
                units.append(value)
                return
            for seg in candidates:
                units.append(seg)

        def flush() -> None:
            if buffer:
                chunk = " ".join(buffer).strip()
                if len(chunk) > 220 and "." in chunk:
                    append_sentence_units(chunk)
                else:
                    units.append(chunk)
                buffer.clear()

        for line in lines:
            if not line:
                flush()
                continue

            is_heading = bool(re.match(r"^(?:\d+(?:\.\d+)*[\).:-]?\s+)?[A-Z0-9][A-Z0-9 ,/&()\-]{4,}$", line))
            is_numbered = bool(re.match(r"^\d+(?:\.\d+)*[\).:-]\s+", line))

            if is_heading or is_numbered:
                flush()
                units.append(line)
                continue

            if len(line) < 80 and not line.endswith((".", ";", ":")):
                buffer.append(line)
                continue

            if buffer:
                buffer.append(line)
                flush()
            else:
                if len(line) > 220 and "." in line:
                    append_sentence_units(line)
                else:
                    units.append(line)

        flush()
        return [unit for unit in units if unit.strip()]

    def _score_rule_hits(self, units: List[str], rules: List[Dict[str, Any]]) -> List[Dict[str, Any]]:
        hits: List[Dict[str, Any]] = []
        for index, unit in enumerate(units):
            lower = unit.lower()
            for rule in rules:
                matched = [keyword for keyword in rule["keywords"] if keyword in lower]
                if not matched:
                    continue
                score = len(matched)
                if any(keyword in lower for keyword in ("must", "shall", "will", "may not", "prohibited")):
                    score += 1
                hits.append(
                    {
                        "category": rule["category"],
                        "title": rule["title"],
                        "snippet": unit[:280] + ("..." if len(unit) > 280 else ""),
                        "score": score,
                        "index": index,
                        "keywords": matched,
                    }
                )
        hits.sort(key=lambda item: (-float(item["score"]), int(item["index"])))
        return hits

    def _unique_by_category(self, hits: List[Dict[str, Any]], limit: int = 8) -> List[Dict[str, Any]]:
        selected: List[Dict[str, Any]] = []
        seen = set()
        for hit in hits:
            category = str(hit.get("category", "")).strip()
            if not category or category in seen:
                continue
            seen.add(category)
            selected.append(hit)
            if len(selected) >= limit:
                break
        return selected

    def _domain_tags(self, legal_hits: List[Dict[str, Any]], startup_hits: List[Dict[str, Any]]) -> List[str]:
        tags = {"legal"}
        if startup_hits:
            tags.add("startup")
        for hit in legal_hits[:6] + startup_hits[:6]:
            tags.add(str(hit.get("category", "")).strip())
        return sorted(tag for tag in tags if tag)

    def _first_sentence(self, text: str) -> str:
        match = re.search(r"(.+?[.!?])(?:\s|$)", text)
        if match:
            return match.group(1).strip()
        first_line = next((line.strip() for line in text.splitlines() if line.strip()), "")
        return first_line[:220]

    def _clause_title(self, text: str) -> str:
        words = re.findall(r"[A-Za-z0-9][A-Za-z0-9'&/\-]*", text)
        if not words:
            return "Clause"
        title = " ".join(words[:8])
        return title[:80].strip() or "Clause"

    def _sentence_score(self, sentence: str, rules: List[Dict[str, Any]]) -> float:
        lower = sentence.lower()
        score = 0.0
        if any(keyword in lower for keyword in ("shall", "must", "will", "may not", "prohibited", "required", "entitled")):
            score += 2.0
        if any(keyword in lower for keyword in ("agreement", "party", "confidential", "liability", "indemnity", "governing law", "jurisdiction", "termination", "assignment", "license", "equity", "investment", "board", "payment", "privacy", "data protection")):
            score += 1.0
        if re.search(r"\b\d{1,2}(?:/\d{1,2})?(?:/\d{2,4})?\b|\b\d{4}\b", sentence):
            score += 0.5
        if any(symbol in sentence for symbol in ("$", "€", "%")):
            score += 0.4
        if len(sentence) > 120:
            score += 0.6
        if len(sentence) < 25:
            score -= 0.4

        for rule in rules:
            matched = [keyword for keyword in rule["keywords"] if keyword in lower]
            if matched:
                score += min(2.5, 0.6 + 0.4 * len(matched))
        return score

    def _unique_by_snippet_and_category(self, hits: List[Dict[str, Any]], limit: int = 8) -> List[Dict[str, Any]]:
        selected: List[Dict[str, Any]] = []
        seen = set()
        for hit in hits:
            category = str(hit.get("category", "")).strip()
            snippet = str(hit.get("snippet", "")).strip().lower()
            fingerprint = (category, snippet)
            if not category or fingerprint in seen:
                continue
            seen.add(fingerprint)
            selected.append(hit)
            if len(selected) >= limit:
                break
        return selected

    def _prompt_text(self, text: str, limit: int = 12000) -> str:
        normalized = self._normalize_text(text)
        if len(normalized) <= limit:
            return normalized
        return normalized[:limit] + "\n[truncated]"

    def _compact_hit(self, hit: Dict[str, Any]) -> Dict[str, Any]:
        return {
            "category": hit.get("category"),
            "title": hit.get("title"),
            "snippet": hit.get("snippet"),
            "score": hit.get("score"),
            "keywords": list(hit.get("keywords", []))[:8],
        }

    def _candidate_bundle(self, legal_hits: List[Dict[str, Any]], startup_hits: List[Dict[str, Any]], max_items: int = 8) -> Dict[str, Any]:
        return {
            "legal_candidates": [self._compact_hit(hit) for hit in legal_hits[:max_items]],
            "startup_candidates": [self._compact_hit(hit) for hit in startup_hits[:max_items]],
        }

    def _strip_json_fences(self, content: str) -> str:
        text = (content or "").strip()
        fenced = re.search(r"```(?:json)?\s*(\{.*\})\s*```", text, re.DOTALL)
        if fenced:
            return fenced.group(1).strip()
        if text.startswith("`") and text.endswith("`") and len(text) > 2:
            inner = text[1:-1].strip()
            if inner.startswith("{") and inner.endswith("}"):
                return inner
        return text

    def _parse_json_object(self, content: str) -> Optional[Dict[str, Any]]:
        text = self._strip_json_fences(content)
        if not text:
            return None
        try:
            parsed = json.loads(text)
        except Exception:
            start = text.find("{")
            end = text.rfind("}")
            if start < 0 or end <= start:
                return None
            try:
                parsed = json.loads(text[start : end + 1])
            except Exception:
                return None
        return parsed if isinstance(parsed, dict) else None

    def _provider_sequence(self) -> List[str]:
        providers: List[str] = []
        openai_ready = self.openai_configured
        ollama_ready = self.ollama_configured
        local_ready = self.local_configured

        if self.ai_provider == "openai":
            providers.append("openai")
            if ollama_ready:
                providers.append("ollama")
            return providers

        if self.ai_provider == "ollama":
            providers.append("ollama")
            if openai_ready:
                providers.append("openai")
            if local_ready:
                providers.append("local")
            return providers

        if self.ai_provider in {"local", "finetuned", "llama_local"}:
            providers.append("local")
            if ollama_ready:
                providers.append("ollama")
            if openai_ready:
                providers.append("openai")
            return providers

        if openai_ready:
            providers.append("openai")
        if ollama_ready:
            providers.append("ollama")
        if local_ready:
            providers.append("local")
        return providers

    def _provider_config(self, provider: str) -> Dict[str, str]:
        if provider == "openai":
            return {
                "model": self.openai_model,
                "base_url": self.openai_base_url,
                "api_key": self.openai_api_key,
            }
        if provider == "ollama":
            return {
                "model": self.ollama_model,
                "base_url": self.ollama_base_url,
                "api_key": self.ollama_api_key,
            }
        if provider == "local":
            return {
                "model": self.local_model_path,
                "base_url": "",
                "api_key": "",
            }
        raise RuntimeError(f"Unsupported AI provider: {provider}")

    def _load_local_model(self, model_path: str):
        if AutoModelForCausalLM is None or AutoTokenizer is None:
            raise RuntimeError(
                "Local LLaMA inference requires transformers. Install transformers and torch."
            )
        path = str(Path(model_path).expanduser().resolve())
        if self.__class__._LOCAL_MODEL is not None and self.__class__._LOCAL_MODEL_PATH == path:
            return self.__class__._LOCAL_MODEL, self.__class__._LOCAL_TOKENIZER

        device_map = "auto"
        dtype = None
        if torch is not None:
            if torch.cuda.is_available():
                dtype = torch.float16
            elif getattr(torch.backends, "mps", None) and torch.backends.mps.is_available():
                dtype = torch.float16
                device_map = {"": "mps"}

        tokenizer = AutoTokenizer.from_pretrained(path)
        if tokenizer.pad_token is None:
            tokenizer.pad_token = tokenizer.eos_token

        load_kwargs: Dict[str, Any] = {"device_map": device_map}
        if dtype is not None:
            load_kwargs["torch_dtype"] = dtype

        model = AutoModelForCausalLM.from_pretrained(path, **load_kwargs)
        model.eval()
        self.__class__._LOCAL_MODEL = model
        self.__class__._LOCAL_TOKENIZER = tokenizer
        self.__class__._LOCAL_MODEL_PATH = path
        return model, tokenizer

    def _invoke_local_model(self, payload: Dict[str, Any], model_path: str) -> Dict[str, Any]:
        model, tokenizer = self._load_local_model(model_path)
        prompt = (
            "You are a legal document analyst. Return ONLY valid JSON with keys: "
            "document_kind, summary, domain_tags, key_clause_indices, startup_signal_indices, "
            "risk_flags, open_questions, answer, evidence.\n"
            f"Input payload:\n{json.dumps(payload, ensure_ascii=True)}\n"
            "JSON output:"
        )

        encoded = tokenizer(prompt, return_tensors="pt")
        if hasattr(model, "device"):
            encoded = {k: v.to(model.device) for k, v in encoded.items()}
        output = model.generate(
            **encoded,
            max_new_tokens=self.local_max_new_tokens,
            do_sample=False,
            pad_token_id=tokenizer.eos_token_id,
            eos_token_id=tokenizer.eos_token_id,
        )
        decoded = tokenizer.decode(output[0], skip_special_tokens=True)
        generated = decoded[len(prompt) :].strip() if decoded.startswith(prompt) else decoded
        parsed = self._parse_json_object(generated)
        if parsed is None:
            raise RuntimeError("Local LLaMA returned non-JSON output")
        return parsed

    def _invoke_ai_for_provider(self, provider: str, payload: Dict[str, Any]) -> Dict[str, Any]:
        config = self._provider_config(provider)
        if provider == "local":
            model_path = config.get("model", "")
            if not model_path:
                raise RuntimeError("Local provider selected but DOCUMENT_INTEL_LOCAL_MODEL_PATH is empty")
            return self._invoke_local_model(payload, model_path)

        base_url = config["base_url"]

        system_prompt = (
            "You are a legal document analyst for startup and commercial agreements. "
            "Return only JSON. Do not invent facts. Prefer the provided candidates, but you may "
            "derive a better summary from the full text. Use concise business language."
        )
        user_prompt = json.dumps(payload, ensure_ascii=True)
        messages = [
            {"role": "system", "content": system_prompt},
            {"role": "user", "content": user_prompt},
        ]

        request_body = {
            "model": config["model"],
            "messages": messages,
            "temperature": 0.2,
        }
        if provider == "openai":
            request_body["response_format"] = {"type": "json_object"}
        headers = {"Content-Type": "application/json"}
        if config["api_key"]:
            headers["Authorization"] = f"Bearer {config['api_key']}"

        response = httpx.post(
            f"{base_url}/chat/completions",
            headers=headers,
            json=request_body,
            timeout=self.ai_timeout,
        )
        response.raise_for_status()
        data = response.json()
        content = (((data.get("choices") or [{}])[0]).get("message") or {}).get("content", "")
        parsed = self._parse_json_object(str(content))
        if parsed is None:
            raise RuntimeError("LLM returned a non-JSON response")
        return parsed

    def _invoke_ai(self, payload: Dict[str, Any]) -> Optional[Dict[str, Any]]:
        providers = self._provider_sequence()
        if not providers:
            raise RuntimeError("No LLM backend configured. Set OpenAI, Ollama, or local model environment variables.")

        last_error: Optional[Exception] = None
        for provider in providers:
            try:
                return self._invoke_ai_for_provider(provider, payload)
            except Exception as exc:
                last_error = exc

        provider_list = ", ".join(providers)
        raise RuntimeError(f"Document intelligence LLM request failed via {provider_list}") from last_error

    def _select_hits_by_index(self, hits: List[Dict[str, Any]], indexes: List[Any], limit: int = 8) -> List[Dict[str, Any]]:
        selected: List[Dict[str, Any]] = []
        used = set()
        for raw_index in indexes or []:
            try:
                index = int(raw_index)
            except Exception:
                continue
            if index < 0 or index >= len(hits) or index in used:
                continue
            used.add(index)
            selected.append(hits[index])
            if len(selected) >= limit:
                break
        return selected

    def _extract_clause_candidates(self, text: str, rules: List[Dict[str, Any]], limit: int = 8) -> List[Dict[str, Any]]:
        nlp = None
        try:
            nlp = self._load_spacy()
        except Exception:
            nlp = None

        units = self._split_units(text)
        hits: List[Dict[str, Any]] = []

        for index, unit in enumerate(units):
            lower = unit.lower()
            matched_keywords: List[str] = []
            matched_rules: List[Dict[str, Any]] = []
            for rule in rules:
                matched = [keyword for keyword in rule["keywords"] if keyword in lower]
                if matched:
                    matched_rules.append(
                        {
                            "rule": rule,
                            "matched": matched,
                        }
                    )
                matched_keywords.extend(matched)

            score = self._sentence_score(unit, rules)
            if re.match(r"^(?:\d+(?:\.\d+)*[\).:-]?\s+)?[A-Z0-9][A-Z0-9 ,/&()\-]{4,}$", unit):
                score += 1.0

            if nlp is not None:
                try:
                    doc = nlp(unit)
                    entity_labels = {ent.label_ for ent in doc.ents}
                    if entity_labels & {"ORG", "PERSON", "GPE", "DATE", "MONEY"}:
                        score += 0.75
                    if len(list(doc.ents)) >= 2:
                        score += 0.25
                except Exception:
                    pass

            if score < 1.0 and not matched_rules:
                continue

            snippet = unit[:280] + ("..." if len(unit) > 280 else "")
            if matched_rules:
                # Emit a candidate per matched rule to improve category recall.
                for rule_hit in matched_rules:
                    rule = rule_hit["rule"]
                    rule_keywords = self._dedupe_text(rule_hit["matched"])
                    hits.append(
                        {
                            "category": rule["category"],
                            "title": rule["title"],
                            "snippet": snippet,
                            "score": round(score + min(1.5, 0.2 * len(rule_keywords)), 3),
                            "index": index,
                            "keywords": rule_keywords,
                        }
                    )
            else:
                hits.append(
                    {
                        "category": "clause",
                        "title": self._clause_title(unit),
                        "snippet": snippet,
                        "score": round(score, 3),
                        "index": index,
                        "keywords": self._dedupe_text(matched_keywords),
                    }
                )

        hits.sort(key=lambda item: (-float(item["score"]), int(item["index"])))
        return self._unique_by_snippet_and_category(hits, limit=limit)

    def _compose_summary(
        self,
        doc_kind: str,
        source_name: Optional[str],
        legal_hits: List[Dict[str, Any]],
        startup_hits: List[Dict[str, Any]],
        text: str,
    ) -> str:
        pieces: List[str] = []
        opener = self._first_sentence(text)
        if opener:
            pieces.append(opener)

        focus_bits: List[str] = []
        for hit in (legal_hits[:3] + startup_hits[:2]):
            title = str(hit.get("title", "")).strip()
            if title and title not in focus_bits:
                focus_bits.append(title)

        if focus_bits:
            pieces.append(f"This appears to be a {doc_kind} centered on {', '.join(focus_bits[:4])}.")
        else:
            pieces.append(f"This appears to be a {doc_kind}.")

        if source_name:
            pieces.append(f"Source: {source_name}.")

        return " ".join(pieces).strip()

    def _summary_from_hits(
        self,
        doc_kind: str,
        legal_hits: List[Dict[str, Any]],
        startup_hits: List[Dict[str, Any]],
        max_points: int = 4,
    ) -> str:
        """Build a concise, clause-aware summary from top extracted hits."""
        points: List[str] = []
        seen = set()

        for hit in legal_hits[: max_points * 2] + startup_hits[: max_points * 2]:
            category = str(hit.get("category", "")).strip().lower()
            title = str(hit.get("title", "Clause")).strip()
            snippet = str(hit.get("snippet", "")).strip()
            if not category or category in seen or not snippet:
                continue
            seen.add(category)
            first_sentence = self._first_sentence(snippet)
            if not first_sentence:
                continue
            points.append(f"{title}: {first_sentence}")
            if len(points) >= max_points:
                break

        if not points:
            return ""

        return f"{doc_kind.capitalize()} summary. " + " ".join(points)

    def _top_summary_terms(self, legal_hits: List[Dict[str, Any]], startup_hits: List[Dict[str, Any]], limit: int = 10) -> List[str]:
        terms: List[str] = []
        for hit in legal_hits[:6] + startup_hits[:6]:
            for kw in hit.get("keywords", []):
                value = str(kw).strip().lower()
                if value and value not in terms:
                    terms.append(value)
                if len(terms) >= limit:
                    return terms
        return terms

    def _summary_keyword_coverage(self, summary: str, terms: List[str]) -> float:
        if not summary or not terms:
            return 0.0
        low = summary.lower()
        matched = sum(1 for term in terms if term in low)
        return matched / max(len(terms), 1)

    def _extractive_summary(self, text: str, top_k: int = 3) -> str:
        """Simple extractive summarizer: score sentences and return top_k sentences.

        Uses the existing _sentence_score to rank sentences so the method is
        lightweight and deterministic (no external packages required).
        """
        if not text:
            return ""
        # split into sentences (simple rule-based split)
        sentences = [s.strip() for s in re.split(r'(?<=[.!?])\s+', text) if s.strip()]
        if not sentences:
            return ""

        rules = self.LEGAL_RULES + self.STARTUP_RULES
        scored = []
        for s in sentences:
            try:
                score = self._sentence_score(s, rules)
            except Exception:
                score = 0.0
            scored.append((score, s))

        scored.sort(key=lambda t: -t[0])
        top = [s for _, s in scored[: max(1, min(top_k, len(scored)))]]
        return " ".join(top).strip()

    def _missing_review_items(self, domain_tags: List[str], legal_hits: List[Dict[str, Any]], startup_hits: List[Dict[str, Any]]) -> List[str]:
        if "legal" not in domain_tags:
            return []

        categories = {str(hit.get("category", "")).strip() for hit in legal_hits + startup_hits}
        missing: List[str] = []
        for category, note in self.MISSING_REVIEW_ITEMS:
            if category not in categories:
                missing.append(note)
        return missing

    def _extract_question_terms(self, question: str) -> List[str]:
        raw_terms = [part.lower() for part in re.findall(r"[A-Za-z][A-Za-z\-/]{2,}", question or "")]
        stop_words = {
            "what",
            "which",
            "that",
            "this",
            "with",
            "from",
            "there",
            "their",
            "about",
            "could",
            "would",
            "should",
            "where",
            "when",
            "does",
            "do",
            "how",
            "many",
            "much",
            "into",
            "under",
            "over",
            "have",
            "has",
            "will",
            "your",
            "our",
            "the",
            "and",
            "for",
            "are",
            "can",
            "who",
        }
        return [term for term in raw_terms if term not in stop_words]

    def _best_evidence(self, text: str, question: str, hits: List[Dict[str, Any]]) -> List[str]:
        question_terms = self._extract_question_terms(question)
        if not question_terms:
            return []

        evidence: List[str] = []
        low_question = question.lower()
        for hit in hits:
            keywords = [str(keyword).lower() for keyword in hit.get("keywords", [])]
            if keywords and any(keyword in low_question for keyword in keywords):
                evidence.append(str(hit.get("snippet", "")))

        if evidence:
            return evidence[:3]

        best_unit = ""
        best_score = 0
        for unit in self._split_units(text):
            low_unit = unit.lower()
            score = sum(1 for term in question_terms if term in low_unit)
            if score > best_score:
                best_score = score
                best_unit = unit

        return [best_unit[:280] + ("..." if len(best_unit) > 280 else "")] if best_unit else []

    def analyze_text(self, text: str, source_name: Optional[str] = None, question: Optional[str] = None) -> Dict[str, Any]:
        normalized = self._normalize_text(text)
        legal_hits = self._extract_clause_candidates(normalized, self.LEGAL_RULES, limit=8)
        startup_hits = self._extract_clause_candidates(normalized, self.STARTUP_RULES, limit=8)

        legal_score = len([hit for hit in legal_hits if hit.get("category") != "clause"]) + (1 if legal_hits else 0)
        startup_score = len([hit for hit in startup_hits if hit.get("category") != "clause"]) + (1 if startup_hits else 0)
        if startup_score >= 3 and legal_score >= 3:
            doc_kind = "startup legal document"
        elif startup_score >= 3:
            doc_kind = "startup document"
        elif legal_score >= 2:
            doc_kind = "legal document"
        else:
            doc_kind = "general document"

        domain_tags = self._domain_tags(legal_hits, startup_hits)
        summary = self._summary_from_hits(doc_kind, legal_hits, startup_hits, max_points=4)
        if not summary:
            summary = self._compose_summary(doc_kind, source_name, legal_hits, startup_hits, normalized)
        risk_flags = [
            f"Potential review area: {hit['title']}" for hit in legal_hits[:4] if hit.get("category") in {"liability", "governing_law", "termination", "assignment"}
        ]
        risk_flags.extend(
            [f"Startup issue to review: {hit['title']}" for hit in startup_hits[:3] if hit.get("category") in {"equity", "governance", "fundraising", "fund_use"}]
        )
        open_questions = self._missing_review_items(domain_tags, legal_hits, startup_hits)
        risk_flags.extend(open_questions[:4])

        answer = None
        evidence: List[str] = []
        if question and question.strip():
            answer, evidence = self.answer_question(normalized, question, legal_hits + startup_hits)

        # compute an extractive summary early for hybrid/analysis uses
        extractive_summary = ""
        try:
            extractive_summary = self._extractive_summary(normalized, top_k=3)
        except Exception:
            extractive_summary = ""

        ai_payload: Dict[str, Any] = {}
        ai_error: Optional[str] = None
        if self.mode not in {"rule_based", "rules", "rule-only"}:
            try:
                ai_payload = self._invoke_ai(
                    {
                        "source_name": source_name,
                        "document_kind_hint": doc_kind,
                        "question": question or "",
                        "document_text": self._prompt_text(normalized),
                        "candidate_clauses": self._candidate_bundle(legal_hits, startup_hits),
                        "extractive_summary": extractive_summary,
                        "must_cover_terms": self._top_summary_terms(legal_hits, startup_hits, limit=10),
                        "output_format": {
                            "document_kind": "short label",
                            "summary": "one or two sentences",
                            "domain_tags": ["short", "tags"],
                            "key_clause_indices": [0],
                            "startup_signal_indices": [0],
                            "risk_flags": ["short notes"],
                            "open_questions": ["short notes"],
                            "answer": "optional answer to the question",
                            "evidence": ["short supporting excerpts"],
                        },
                    }
                )
            except Exception as exc:
                ai_error = str(exc)
                ai_payload = {}

        ai_doc_kind = str(ai_payload.get("document_kind") or doc_kind).strip()
        if ai_doc_kind:
            doc_kind = ai_doc_kind

        ai_summary = str(ai_payload.get("summary") or "").strip()
        # Decide final summary: prefer AI summary when present, otherwise use extractive.
        summary_sources: List[str] = []
        must_cover_terms = self._top_summary_terms(legal_hits, startup_hits, limit=10)
        if ai_summary:
            ai_cov = self._summary_keyword_coverage(ai_summary, must_cover_terms)
            ex_cov = self._summary_keyword_coverage(extractive_summary, must_cover_terms)
            # Guardrail: if AI drops too many key terms, keep extractive summary.
            if extractive_summary and ai_cov + 0.15 < ex_cov:
                summary = extractive_summary
                summary_sources.append("extractive_guardrail")
            else:
                summary = ai_summary
                summary_sources.append("ai")
        if extractive_summary:
            # always include the extractive summary as provenance
            summary_sources.append("extractive")
            # if there was no ai summary, use extractive as the main summary
            if not ai_summary:
                summary = extractive_summary

        ai_tags = [str(tag).strip() for tag in ai_payload.get("domain_tags", []) if str(tag).strip()]
        if ai_tags:
            domain_tags = self._dedupe_text(domain_tags + ai_tags)

        ai_risks = [str(item).strip() for item in ai_payload.get("risk_flags", []) if str(item).strip()]
        if ai_risks:
            risk_flags = self._dedupe_text(ai_risks + risk_flags)

        ai_questions = [str(item).strip() for item in ai_payload.get("open_questions", []) if str(item).strip()]
        if ai_questions:
            open_questions = self._dedupe_text(ai_questions + open_questions)

        ai_answer = str(ai_payload.get("answer") or "").strip()
        if ai_answer:
            answer = ai_answer

        ai_evidence = [str(item).strip() for item in ai_payload.get("evidence", []) if str(item).strip()]
        if ai_evidence:
            evidence = self._dedupe_text(ai_evidence)[:3]

        key_clause_indices = ai_payload.get("key_clause_indices", [])
        startup_signal_indices = ai_payload.get("startup_signal_indices", [])
        selected_legal_hits = self._select_hits_by_index(legal_hits, key_clause_indices, limit=8)
        selected_startup_hits = self._select_hits_by_index(startup_hits, startup_signal_indices, limit=8)
        if selected_legal_hits:
            legal_hits = selected_legal_hits
        if selected_startup_hits:
            startup_hits = selected_startup_hits

        return {
            "source_name": source_name,
            "document_kind": doc_kind,
            "summary": summary,
            "summary_sources": summary_sources,
            "extractive_summary": extractive_summary,
            "ai_error": ai_error,
            "domain_tags": domain_tags,
            "key_clauses": legal_hits,
            "startup_signals": startup_hits,
            "risk_flags": self._dedupe_text(risk_flags),
            "open_questions": self._dedupe_text(open_questions),
            "answer": answer,
            "evidence": evidence,
        }

    def _dedupe_text(self, items: List[str]) -> List[str]:
        seen = set()
        output: List[str] = []
        for item in items:
            value = str(item).strip()
            if not value:
                continue
            low = value.lower()
            if low in seen:
                continue
            seen.add(low)
            output.append(value)
        return output

    def load_text_from_path(self, file_path: str) -> str:
        path = Path(file_path)
        suffix = path.suffix.lower()

        if suffix == ".docx":
            return self._load_docx_text(str(path))

        if suffix in {".txt", ".md"}:
            return self._normalize_text(path.read_text(encoding="utf-8"))

        if suffix in self.IMAGE_EXTENSIONS:
            return self._load_image_text(str(path))

        raise ValueError(f"Unsupported text document type: {suffix or 'unknown'}")

    def answer_question(
        self,
        text: str,
        question: str,
        hits: Optional[List[Dict[str, Any]]] = None,
    ) -> tuple[str, List[str]]:
        units = self._split_units(text)
        matched_hits = hits or self._extract_clause_candidates(text, self.LEGAL_RULES + self.STARTUP_RULES, limit=10)
        evidence = self._best_evidence(text, question, matched_hits)

        if evidence:
            answer = f"Most likely relevant clause: {evidence[0]}"
            if len(evidence) > 1:
                answer += " Additional related text is available in the cited snippets."
            return answer, evidence

        question_terms = self._extract_question_terms(question)
        if not question_terms:
            return "I could not map that question to a specific clause in the document.", []

        for unit in units:
            low_unit = unit.lower()
            if any(term in low_unit for term in question_terms):
                snippet = unit[:320] + ("..." if len(unit) > 320 else "")
                return f"Most likely relevant clause: {snippet}", [snippet]

        return (
            "I could not find a precise clause match, but the document may need a manual review of the related section.",
            [],
        )
