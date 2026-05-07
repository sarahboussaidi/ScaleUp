import re
from pathlib import Path

from docx import Document

class TemplateDocumentGenerator:
    def __init__(self):
        pass

    # Replace text inside paragraphs
    def replace_text(self, paragraph, data, sequence_state):
        # Replace longer tokens first so wrappers like
        # [OPTION 2: ... [insert] ...] are processed before inner [insert] tokens.
        ordered_keys = sorted(data.keys(), key=len, reverse=True)

        for key in ordered_keys:
            value = data[key]
            if key not in paragraph.text:
                continue

            if isinstance(value, (list, tuple)):
                # Replace repeated placeholders one-by-one in document order.
                while key in paragraph.text and sequence_state[key] < len(value):
                    replacement = str(value[sequence_state[key]])
                    paragraph.text = paragraph.text.replace(key, replacement, 1)
                    sequence_state[key] += 1
            else:
                paragraph.text = paragraph.text.replace(key, str(value))

    # Replace everywhere in document
    def replace_all(self, doc, data, sequence_state):
        # Paragraphs
        for paragraph in doc.paragraphs:
            self.replace_text(paragraph, data, sequence_state)

        # Tables
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    for paragraph in cell.paragraphs:
                        self.replace_text(paragraph, data, sequence_state)

    def _collect_document_text(self, doc):
        parts = []

        for paragraph in doc.paragraphs:
            parts.append(paragraph.text)

        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    for paragraph in cell.paragraphs:
                        parts.append(paragraph.text)

        return "\n".join(parts)

    def _find_bracket_placeholders(self, doc):
        text = self._collect_document_text(doc)
        return sorted(set(re.findall(r"\[[^\]]+\]", text)))

    def _strip_all_square_brackets(self, doc):
        for paragraph in doc.paragraphs:
            paragraph.text = paragraph.text.replace("[", "").replace("]", "")

        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    for paragraph in cell.paragraphs:
                        paragraph.text = paragraph.text.replace("[", "").replace("]", "")

    # Generate final document
    def generate(self, template_path, output_path, data, strict=True, strip_bracket_artifacts=True):
        template_file = Path(template_path)
        output_file = Path(output_path)

        if not template_file.exists():
            raise FileNotFoundError(f"Template not found: {template_path}")

        doc = Document(template_path)

        sequence_state = {
            key: 0 for key, value in data.items() if isinstance(value, (list, tuple))
        }

        self.replace_all(doc, data, sequence_state)

        unresolved = self._find_bracket_placeholders(doc)
        unused_sequence_values = [
            key
            for key, idx in sequence_state.items()
            if idx < len(data[key])
        ]

        if strict and (unresolved or unused_sequence_values):
            details = []

            if unresolved:
                details.append(
                    "Unreplaced placeholders: " + ", ".join(unresolved)
                )

            if unused_sequence_values:
                details.append(
                    "Unused sequence values for placeholders: "
                    + ", ".join(unused_sequence_values)
                )

            raise ValueError("Template replacement incomplete. " + " | ".join(details))

        # Safety pass: remove any lingering square-bracket artifacts from the final text.
        # This keeps outputs user-friendly when templates contain non-standard bracket remnants.
        if strip_bracket_artifacts:
            full_text = self._collect_document_text(doc)
            if "[" in full_text or "]" in full_text:
                self._strip_all_square_brackets(doc)

        output_file.parent.mkdir(parents=True, exist_ok=True)
        doc.save(output_path)

        return output_path