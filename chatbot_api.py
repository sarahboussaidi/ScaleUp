import re
import uuid
import json
from functools import lru_cache
from collections import Counter
from datetime import datetime
from pathlib import Path
from typing import Any, Dict, List, Optional

from fastapi import FastAPI, HTTPException, UploadFile, File
from fastapi.responses import FileResponse
from fastapi.middleware.cors import CORSMiddleware
from fastapi.staticfiles import StaticFiles
from pydantic import BaseModel, Field
from docx import Document
try:
    from dotenv import load_dotenv
except Exception:
    def load_dotenv(*_args: Any, **_kwargs: Any) -> bool:
        return False

from intelligent_nda_filler import IntelligentNDAFiller
from legal_document_intelligence import LegalDocumentIntelligence
from objective_registry import (
    get_selected_model_path,
    get_selected_value,
    load_objective_report,
    load_objective_selections,
    save_objective_report,
    save_objective_selections,
)
from template_generator import TemplateDocumentGenerator
from ocr.document_analyzer import DocumentAnalyzer
from ocr.multi_ocr_engine import MultiOCREngine

import torch
import torch.nn as nn
from torchvision import transforms
from torchvision import models
from PIL import Image
import numpy as np
import io
from starlette.responses import JSONResponse
from xai.gradcam import GradCAM, overlay_cam_on_image, cam_stats, draw_focus_box
from xai.ocr_explain import ocr_confidence_heatmap, ocr_explanation_text, draw_ocr_boxes
from xai.detection_explain import detection_explanation_text, detection_focus_regions
from xai.summarization_explain import (
    clause_explanation,
    summarization_explanation_text,
    clause_confidence_score,
    highlight_keywords_in_clause,
)


load_dotenv(Path(__file__).resolve().parent / ".env")

try:
    from ultralytics import YOLO
except Exception:
    YOLO = None


class StartResponse(BaseModel):
    session_id: str
    message: str
    templates: List[str]


class ChatMessageRequest(BaseModel):
    session_id: str
    message: str = Field(min_length=1)


class ChatMessageResponse(BaseModel):
    session_id: str
    stage: str
    reply: str
    options: Optional[List[str]] = None
    data_preview: Optional[Dict[str, Any]] = None
    can_generate: bool = False


class GenerateRequest(BaseModel):
    session_id: str


class GenerateResponse(BaseModel):
    session_id: str
    document_path: str
    report_path: Optional[str] = None
    warnings: List[str]


class OCRAnalyzeRequest(BaseModel):
    image_path: str
    save_annotated: bool = True


class OCRAnalyzeResponse(BaseModel):
    image_path: str
    text: str
    entities: List[Dict[str, str]]
    keyword_hits: List[str]
    ocr_results: List[Dict[str, Any]]
    annotated_image_path: Optional[str] = None
    annotated_image_url: Optional[str] = None


class DocumentIntelligenceRequest(BaseModel):
    text: Optional[str] = None
    document_path: Optional[str] = None
    question: Optional[str] = None
    source_name: Optional[str] = None


class DocumentIntelligenceResponse(BaseModel):
    source_name: Optional[str] = None
    document_kind: str
    summary: str
    domain_tags: List[str]
    key_clauses: List[Dict[str, Any]]
    startup_signals: List[Dict[str, Any]]
    risk_flags: List[str]
    open_questions: List[str]
    answer: Optional[str] = None
    evidence: List[str] = Field(default_factory=list)


class SessionState(BaseModel):
    session_id: str
    stage: str
    template_name: Optional[str] = None
    template_path: Optional[str] = None
    is_nda_like: bool = False
    placeholders: Dict[str, int] = Field(default_factory=dict)
    question_queue: List[Dict[str, Any]] = Field(default_factory=list)
    question_index: int = 0
    pending_missing_queue: List[Dict[str, Any]] = Field(default_factory=list)
    pending_missing_index: int = 0
    data: Dict[str, Any] = Field(default_factory=dict)
    report: Dict[str, Any] = Field(default_factory=dict)


app = FastAPI(title="Legal Chatbot API", version="1.0.0")
app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"],
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)

# --- Signature Prediction Endpoint (after app is defined) ---

# --- Signature Prediction Endpoint (must be after app is defined) ---

# Place this after app = FastAPI(...)


class SignatureCNN(nn.Module):
    def __init__(self) -> None:
        super().__init__()
        self.conv = nn.Sequential(
            nn.Conv2d(1, 16, 3, 1, 1),
            nn.ReLU(),
            nn.MaxPool2d(2),
            nn.Conv2d(16, 32, 3, 1, 1),
            nn.ReLU(),
            nn.MaxPool2d(2),
            nn.Conv2d(32, 64, 3, 1, 1),
            nn.ReLU(),
            nn.MaxPool2d(2),
        )
        self.fc = nn.Sequential(
            nn.Flatten(),
            nn.Linear(64 * 16 * 32, 128),
            nn.ReLU(),
            nn.Linear(128, 2),
        )

    def forward(self, x: torch.Tensor) -> torch.Tensor:
        return self.fc(self.conv(x))


def _rgb_224_transform() -> transforms.Compose:
    return transforms.Compose(
        [
            transforms.Resize((224, 224)),
            transforms.ToTensor(),
            transforms.Normalize([0.5, 0.5, 0.5], [0.5, 0.5, 0.5]),
        ]
    )


def _gray_128x256_transform() -> transforms.Compose:
    return transforms.Compose(
        [
            transforms.Grayscale(num_output_channels=1),
            transforms.Resize((128, 256)),
            transforms.ToTensor(),
            transforms.Normalize([0.5], [0.5]),
        ]
    )


def _selection_cache_key(objective: str) -> str:
    return json.dumps(load_objective_selections().get(objective, {}), sort_keys=True)


def _load_scripted_model(model_path: Path) -> tuple[torch.nn.Module, transforms.Compose]:
    model = torch.jit.load(str(model_path), map_location="cpu")
    model.eval()
    if "cnn" in model_path.name.lower():
        return model, _gray_128x256_transform()
    return model, _rgb_224_transform()


def _load_pth_model(model_path: Path) -> tuple[torch.nn.Module, transforms.Compose]:
    state_dict = torch.load(str(model_path), map_location="cpu")
    model_name = model_path.name.lower()

    if "vit" in model_name:
        vit = models.vit_b_16(weights=None)
        vit.heads.head = nn.Linear(vit.heads.head.in_features, 2)
        vit.load_state_dict(state_dict)
        vit.eval()
        return vit, _rgb_224_transform()

    if "resnet" in model_name:
        resnet = models.resnet18(weights=None)
        resnet.fc = nn.Linear(resnet.fc.in_features, 2)
        resnet.load_state_dict(state_dict)
        resnet.eval()
        return resnet, _rgb_224_transform()

    cnn = SignatureCNN()
    cnn.load_state_dict(state_dict)
    cnn.eval()
    return cnn, _gray_128x256_transform()


@lru_cache(maxsize=4)
def _load_signature_model(selection_key: str = "") -> tuple[torch.nn.Module, transforms.Compose, Path]:
    scripted_candidates = [
        Path("signature_best_script.pt"),
        Path("signature_resnet18_script.pt"),
        Path("signature_cnn_script.pt"),
    ]
    pth_candidates = [
        Path("best_signature_vit.pth"),
        Path("best_signature_resnet18.pth"),
        Path("best_signature_cnn.pth"),
        Path("signature_classifier.pth"),
    ]

    selected_path = get_selected_model_path("signature_classification", scripted_candidates + pth_candidates)
    if selected_path is not None:
        if selected_path.suffix.lower() == ".pt":
            model, transform = _load_scripted_model(selected_path)
        else:
            model, transform = _load_pth_model(selected_path)
        return model, transform, selected_path

    for model_path in scripted_candidates:
        if model_path.exists():
            model, transform = _load_scripted_model(model_path)
            return model, transform, model_path

    for model_path in pth_candidates:
        if model_path.exists():
            model, transform = _load_pth_model(model_path)
            return model, transform, model_path

    raise HTTPException(
        status_code=500,
        detail=(
            "No signature model found. Expected one of: signature_best_script.pt, "
            "signature_resnet18_script.pt, signature_cnn_script.pt, best_signature_vit.pth, "
            "best_signature_resnet18.pth, best_signature_cnn.pth, or signature_classifier.pth"
        ),
    )


@lru_cache(maxsize=4)
def _load_signature_detector(selection_key: str = "") -> tuple[Any, Path]:
    if YOLO is None:
        raise HTTPException(
            status_code=500,
            detail="Ultralytics is not installed. Install requirements and retry.",
        )

    candidates = [
        Path("runs/detect/runs/detect/signature_yolo_clean/weights/best.pt"),
        Path("runs/detect/signature_yolo_clean/weights/best.pt"),
        Path("runs/detect/signature_yolo_results/yolov8_signature/weights/best.pt"),
        Path("data2/chekpoint_last.pt"),
    ]

    selected_path = get_selected_model_path("signature_detection", candidates)
    if selected_path is not None:
        return YOLO(str(selected_path)), selected_path

    for model_path in candidates:
        if model_path.exists():
            return YOLO(str(model_path)), model_path

    raise HTTPException(
        status_code=500,
        detail=(
            "No YOLO signature detector found. Train first and ensure one of these exists: "
            "runs/detect/runs/detect/signature_yolo_clean/weights/best.pt, "
            "runs/detect/signature_yolo_clean/weights/best.pt, "
            "runs/detect/signature_yolo_results/yolov8_signature/weights/best.pt, "
            "or data2/chekpoint_last.pt"
        ),
    )

@app.post("/api/signature/predict")
async def predict_signature_api(file: UploadFile = File(...)):
    """Predict if a signature is real or forged using the trained model."""
    model, transform, model_path = _load_signature_model(_selection_cache_key("signature_classification"))

    contents = await file.read()
    img = Image.open(io.BytesIO(contents)).convert('RGB')
    img_tensor = transform(img).unsqueeze(0)
    with torch.no_grad():
        out = model(img_tensor)
        pred = out.argmax(1).item()
    # Dataset mapping: {'fake': 0, 'real': 1}
    label = 'fake' if pred == 0 else 'real'
    return JSONResponse({"prediction": label, "model": model_path.name})


@app.post("/api/signature/explain")
async def explain_signature_api(file: UploadFile = File(...)):
    """Return an explainable prediction with Grad-CAM artifacts and metrics."""
    model, transform, model_path = _load_signature_model(_selection_cache_key("signature_classification"))

    contents = await file.read()
    img = Image.open(io.BytesIO(contents)).convert('RGB')
    img_tensor = transform(img).unsqueeze(0)

    # If the model is TorchScript, try to find a .pth fallback because hooks need nn.Module
    try:
        is_scripted = isinstance(model, torch.jit.ScriptModule) or ("script" in model_path.name.lower())
    except Exception:
        is_scripted = ("script" in model_path.name.lower())

    if is_scripted:
        # look for common .pth variants
        pth_candidates = [
            PROJECT_ROOT / "best_signature_vit.pth",
            PROJECT_ROOT / "best_signature_resnet18.pth",
            PROJECT_ROOT / "best_signature_cnn.pth",
            PROJECT_ROOT / "signature_classifier.pth",
        ]
        fallback = None
        for p in pth_candidates:
            if p.exists():
                fallback = p
                break
        if fallback:
            model, transform = _load_pth_model(fallback)
            model_path = fallback
            # Re-apply the new transform to the image since the fallback model may require different dimensions
            img_tensor = transform(img).unsqueeze(0)
        else:
            raise HTTPException(
                status_code=500,
                detail=(
                    "Explainability requires a non-scripted model (.pth state_dict). "
                    "No fallback .pth model found."
                ),
            )

    model = model.to("cpu")
    model.eval()

    # prediction
    with torch.no_grad():
        out = model(img_tensor)
        probs = torch.softmax(out, dim=1)[0].cpu().numpy().tolist()
        pred = int(out.argmax(1).item())
    label = 'fake' if pred == 0 else 'real'
    pred_conf = float(probs[pred])

    # generate grad-cam
    try:
        cammer = GradCAM(model)
        cam = cammer.generate_cam(img_tensor, class_idx=pred)
        overlay = overlay_cam_on_image(img, cam, alpha=0.5)
        stats = cam_stats(cam, threshold_quantile=0.85)
        focus_overlay = draw_focus_box(overlay, stats["focus_box_norm_xyxy"], color=(255, 255, 255), width=2)
    except Exception as exc:
        raise HTTPException(status_code=500, detail=f"Grad-CAM generation failed: {exc}") from exc

    stamp = datetime.utcnow().strftime("%Y%m%d_%H%M%S")
    out_name = f"signature_explain_{stamp}_{uuid.uuid4().hex[:8]}.png"
    out_path = GENERATED_DIR / out_name
    focus_overlay.save(out_path)

    saliency_pct = round(stats["salient_area_ratio"] * 100.0, 1)
    cx, cy = stats["centroid_norm_xy"]
    x_region = "left" if cx < 0.4 else "right" if cx > 0.6 else "center"
    y_region = "top" if cy < 0.4 else "bottom" if cy > 0.6 else "middle"
    explanation_text = (
        f"Predicted '{label}' with {pred_conf:.1%} confidence. "
        f"The model focused on about {saliency_pct}% of the signature area, "
        f"mainly in the {y_region}-{x_region} region."
    )

    return JSONResponse(
        {
            "model": model_path.name,
            "prediction": label,
            "prediction_confidence": round(pred_conf, 4),
            "class_probabilities": {
                "fake": round(float(probs[0]), 4),
                "real": round(float(probs[1]), 4),
            },
            "xai_method": "gradcam",
            "xai_stats": stats,
            "explanation_text": explanation_text,
            "annotated_image_url": _generated_url(str(out_path)),
        }
    )


@app.post("/api/signature/detect")
async def detect_signature_api(
    file: UploadFile = File(...),
    conf: float = 0.25,
    iou: float = 0.45,
):
    """Detect signature bounding boxes with YOLO."""
    model, model_path = _load_signature_detector(_selection_cache_key("signature_detection"))

    contents = await file.read()
    image = Image.open(io.BytesIO(contents)).convert("RGB")
    image_np = np.array(image)

    results = model.predict(source=image_np, conf=conf, iou=iou, verbose=False)
    if not results:
        return JSONResponse(
            {
                "model": model_path.name,
                "count": 0,
                "detections": [],
                "annotated_image_url": None,
            }
        )

    result = results[0]
    detections: List[Dict[str, Any]] = []
    names = result.names or {}

    if result.boxes is not None and len(result.boxes) > 0:
        xyxy_list = result.boxes.xyxy.cpu().tolist()
        conf_list = result.boxes.conf.cpu().tolist()
        cls_list = result.boxes.cls.cpu().tolist()

        for xyxy, score, cls_idx in zip(xyxy_list, conf_list, cls_list):
            idx = int(cls_idx)
            detections.append(
                {
                    "class_id": idx,
                    "class_name": str(names.get(idx, idx)),
                    "confidence": round(float(score), 4),
                    "box_xyxy": [round(float(v), 2) for v in xyxy],
                }
            )

    stamp = datetime.utcnow().strftime("%Y%m%d_%H%M%S")
    annotated_name = f"signature_detect_{stamp}_{uuid.uuid4().hex[:8]}.png"
    annotated_path = GENERATED_DIR / annotated_name

    plotted_bgr = result.plot()
    plotted_rgb = plotted_bgr[:, :, ::-1]
    Image.fromarray(plotted_rgb).save(annotated_path)

    return JSONResponse(
        {
            "model": model_path.name,
            "count": len(detections),
            "detections": detections,
            "annotated_image_url": _generated_url(str(annotated_path)),
        }
    )


@app.post("/api/signature/detect-explain")
async def detect_signature_explain_api(file: UploadFile = File(...)):
    """Explain signature detections with confidence and focus regions."""
    model, model_path = _load_signature_detector(_selection_cache_key("signature_detection"))

    contents = await file.read()
    image = Image.open(io.BytesIO(contents)).convert("RGB")
    image_np = np.array(image)

    results = model.predict(source=image_np, conf=0.25, iou=0.45, verbose=False)
    detections: List[Dict[str, Any]] = []

    if results and len(results) > 0:
        result = results[0]
        names = result.names or {}
        if result.boxes is not None and len(result.boxes) > 0:
            xyxy_list = result.boxes.xyxy.cpu().tolist()
            conf_list = result.boxes.conf.cpu().tolist()
            cls_list = result.boxes.cls.cpu().tolist()
            h, w = image_np.shape[:2]

            for xyxy, score, cls_idx in zip(xyxy_list, conf_list, cls_list):
                idx = int(cls_idx)
                x1, y1, x2, y2 = [float(v) for v in xyxy]
                detections.append({
                    "class_id": idx,
                    "class_name": str(names.get(idx, idx)),
                    "confidence": round(float(score), 4),
                    "box_xyxy": [round(v / d, 4) for v, d in zip([x1, y1, x2, y2], [w, h, w, h])],
                })

    explanation = detection_explanation_text(detections)
    focus_regions = detection_focus_regions(detections)

    # Save annotated detection image for UI/inspection
    annotated_url = None
    try:
        if results and len(results) > 0:
            stamp = datetime.utcnow().strftime("%Y%m%d_%H%M%S")
            annotated_name = f"signature_detect_explain_{stamp}_{uuid.uuid4().hex[:8]}.png"
            annotated_path = GENERATED_DIR / annotated_name
            plotted_bgr = results[0].plot()
            plotted_rgb = plotted_bgr[:, :, ::-1]
            Image.fromarray(plotted_rgb).save(annotated_path)
            annotated_url = _generated_url(str(annotated_path))
    except Exception:
        annotated_url = None

    return JSONResponse({
        "model": model_path.name,
        "count": len(detections),
        "detections": detections,
        "xai_method": "yolo_confidence",
        "explanation_text": explanation,
        "focus_regions": focus_regions,
        "annotated_image_url": annotated_url,
    })


@app.post("/api/ocr/analyze-explain")
async def ocr_analyze_explain_api(file: UploadFile = File(...)):
    """Explain OCR results with confidence-based highlighting."""
    contents = await file.read()
    analyzer = DocumentAnalyzer()
    ext = Path(file.filename or "upload.png").suffix.lower() or ".png"
    tmp_dir = GENERATED_DIR / "ocr_uploads"
    tmp_dir.mkdir(parents=True, exist_ok=True)
    stamp = datetime.utcnow().strftime("%Y%m%d_%H%M%S_%f")
    image_path = tmp_dir / f"ocr_explain_{stamp}{ext}"
    image_path.write_bytes(contents)

    annotated_path = GENERATED_DIR / f"ocr_explain_{stamp}_{uuid.uuid4().hex[:8]}.png"

    try:
        result = analyzer.analyze(str(image_path), save_annotated_to=str(annotated_path))
        ocr_results = list(result.get("ocr_results", []))
    except Exception as exc:
        raise HTTPException(status_code=500, detail=f"OCR analysis failed: {exc}") from exc

    explanation = ocr_explanation_text(ocr_results)

    return JSONResponse({
        "xai_method": "ocr_confidence",
        "text_regions_count": len(ocr_results),
        "average_confidence": round(np.mean([r.get('conf', 0) for r in ocr_results]), 4),
        "explanation_text": explanation,
        "annotated_image_url": _generated_url(result.get("annotated_image_path") or str(annotated_path)),
        "ocr_results": ocr_results,
    })


@app.post("/api/document/summarization-explain")
async def summarization_explain_api(request: DocumentIntelligenceRequest):
    """Explain document summarization with clause confidence."""
    try:
        intel = LegalDocumentIntelligence()
        doc_text = None
        annotated_url = None
        if request.document_path:
            p = Path(str(request.document_path))
            suffix = p.suffix.lower()
            if suffix in (".png", ".jpg", ".jpeg", ".tiff", ".bmp") and p.exists():
                # OCR image first so summarization can run on extracted text.
                stamp = datetime.utcnow().strftime("%Y%m%d_%H%M%S")
                out_name = f"summarization_explain_{stamp}_{uuid.uuid4().hex[:8]}.png"
                annotated_path = GENERATED_DIR / out_name
                ocr_res = _ocr_with_selected_engine(p, save_annotated_to=annotated_path)
                doc_text = str(ocr_res.get("text", "")).strip()
                if ocr_res.get("annotated_image_path"):
                    annotated_url = _generated_url(ocr_res.get("annotated_image_path"))
                else:
                    annotated_url = _generated_url(str(annotated_path))
            elif suffix == ".docx" and p.exists():
                doc = Document(str(p))
                doc_text = "\n".join([para.text for para in doc.paragraphs])
            elif p.exists():
                doc_text = p.read_text(encoding="utf-8", errors="ignore")
        elif request.text:
            doc_text = request.text

        if not doc_text:
            raise HTTPException(status_code=400, detail="No document text provided.")

        result = intel.analyze_text(doc_text, question=request.question)

        clauses = result.get("key_clauses", [])
        for clause in clauses:
            category = clause.get("category", "unknown")
            # fallback keywords from the rule-set for this category
            category_keywords = [
                kw
                for rule in intel.LEGAL_RULES
                if rule["category"] == category
                for kw in rule["keywords"]
            ]
            text_snippet = clause.get("relevant_text") or clause.get("snippet", "")
            # compute a data-driven confidence instead of a fixed value
            conf = float(clause_confidence_score(text_snippet, clause.get("keywords", category_keywords)))
            clause["xai_confidence"] = round(conf, 4)
            clause["xai_explanation"] = clause_explanation(category, text_snippet, clause.get("keywords", category_keywords), conf)
            clause["highlighted_snippet"] = highlight_keywords_in_clause(text_snippet, clause.get("keywords", category_keywords))
        explanation = summarization_explanation_text(clauses, len(result.get("key_clauses", [])))

        resp = {
            "xai_method": "clause_confidence",
            "document_kind": result.get("document_kind"),
            "summary": result.get("summary"),
            "summary_sources": result.get("summary_sources", []),
            "extractive_summary": result.get("extractive_summary", ""),
            "explanation_text": explanation,
            "key_clauses_count": len(clauses),
            "key_clauses": clauses[:5],
        }
        if annotated_url:
            resp["annotated_image_url"] = annotated_url

        return JSONResponse(resp)
    except Exception as exc:
        raise HTTPException(status_code=500, detail=f"Summarization analysis failed: {exc}") from exc


@app.post("/api/document/summarization-explain-upload")
async def summarization_explain_upload_api(file: UploadFile = File(...), question: Optional[str] = None):
    """Explain summarization from an uploaded file (image/txt/docx)."""
    try:
        intel = LegalDocumentIntelligence()
        name = (file.filename or "upload").lower()
        suffix = Path(name).suffix.lower()
        contents = await file.read()

        doc_text = ""
        annotated_url = None

        if suffix in (".png", ".jpg", ".jpeg", ".tiff", ".bmp"):
            tmp_name = f"sum_upload_{uuid.uuid4().hex[:8]}{suffix or '.png'}"
            tmp_path = GENERATED_DIR / tmp_name
            tmp_path.write_bytes(contents)

            stamp = datetime.utcnow().strftime("%Y%m%d_%H%M%S")
            out_name = f"summarization_explain_{stamp}_{uuid.uuid4().hex[:8]}.png"
            annotated_path = GENERATED_DIR / out_name
            ocr_res = _ocr_with_selected_engine(tmp_path, save_annotated_to=annotated_path)
            doc_text = str(ocr_res.get("text", "")).strip()
            if ocr_res.get("annotated_image_path"):
                annotated_url = _generated_url(ocr_res.get("annotated_image_path"))
            else:
                annotated_url = _generated_url(str(annotated_path))
        elif suffix == ".docx":
            doc = Document(io.BytesIO(contents))
            doc_text = "\n".join([para.text for para in doc.paragraphs])
        else:
            doc_text = contents.decode("utf-8", errors="ignore")

        if not doc_text.strip():
            raise HTTPException(status_code=400, detail="No document text extracted from uploaded file.")

        result = intel.analyze_text(doc_text, question=question)
        clauses = result.get("key_clauses", [])
        for clause in clauses:
            category = clause.get("category", "unknown")
            category_keywords = [
                kw
                for rule in intel.LEGAL_RULES
                if rule["category"] == category
                for kw in rule["keywords"]
            ]
            text_snippet = clause.get("relevant_text") or clause.get("snippet", "")
            conf = float(clause_confidence_score(text_snippet, clause.get("keywords", category_keywords)))
            clause["xai_confidence"] = round(conf, 4)
            clause["xai_explanation"] = clause_explanation(category, text_snippet, clause.get("keywords", category_keywords), conf)
            clause["highlighted_snippet"] = highlight_keywords_in_clause(text_snippet, clause.get("keywords", category_keywords))

        explanation = summarization_explanation_text(clauses, len(result.get("key_clauses", [])))
        resp = {
            "xai_method": "clause_confidence",
            "document_kind": result.get("document_kind"),
            "summary": result.get("summary"),
            "summary_sources": result.get("summary_sources", []),
            "extractive_summary": result.get("extractive_summary", ""),
            "explanation_text": explanation,
            "key_clauses_count": len(clauses),
            "key_clauses": clauses[:5],
        }
        if annotated_url:
            resp["annotated_image_url"] = annotated_url

        return JSONResponse(resp)
    except Exception as exc:
        raise HTTPException(status_code=500, detail=f"Summarization upload analysis failed: {exc}") from exc


# Multi-OCR comparison endpoints
_multi_ocr = None

def get_multi_ocr_engine():
    """Get or initialize the multi-OCR engine."""
    global _multi_ocr
    selected_engine = str(get_selected_value("ocr", "primary_engine", "easyocr") or "easyocr").strip().lower()
    if _multi_ocr is None:
        _multi_ocr = MultiOCREngine(primary_engine=selected_engine)
    elif _multi_ocr.primary_engine != selected_engine:
        _multi_ocr.set_primary_engine(selected_engine)
    return _multi_ocr


@app.post("/api/ocr/compare-all")
async def ocr_compare_all(file: UploadFile = File(...)):
    """Compare all available OCR engines on the same image."""
    contents = await file.read()
    img_pil = Image.open(io.BytesIO(contents)).convert("RGB")
    img_cv2 = np.array(img_pil)
    
    engine = get_multi_ocr_engine()
    
    try:
        results = engine.extract_all(img_cv2)
    except Exception as exc:
        raise HTTPException(status_code=500, detail=f"OCR comparison failed: {exc}") from exc
    
    # Format results with stats
    comparison = {}
    for eng_name, extracts in results.items():
        if isinstance(extracts, dict) and "error" in extracts:
            comparison[eng_name] = {"error": extracts["error"], "results": []}
        else:
            avg_conf = np.mean([r.get('confidence', 0) for r in extracts]) if extracts else 0
            comparison[eng_name] = {
                "text_regions": len(extracts),
                "average_confidence": round(float(avg_conf), 4),
                "results": extracts[:20]  # First 20 results
            }
    
    return JSONResponse({
        "timestamp": datetime.utcnow().isoformat(),
        "engines": list(comparison.keys()),
        "comparison": comparison,
        "primary_engine": engine.primary_engine,
    })


@app.post("/api/ocr/set-primary")
async def ocr_set_primary(engine: str):
    """Switch to a different primary OCR engine."""
    try:
        multi_ocr = get_multi_ocr_engine()
        multi_ocr.set_primary_engine(engine)
        return JSONResponse({
            "primary_engine": engine,
            "message": f"Switched to {engine}"
        })
    except Exception as exc:
        raise HTTPException(status_code=400, detail=str(exc))


@app.get("/api/ocr/accuracy-stats")
async def ocr_accuracy_stats():
    """Get accuracy statistics for all engines."""
    engine = get_multi_ocr_engine()
    stats = engine.get_accuracy_stats()
    
    return JSONResponse({
        "stats": {k: {
            "count": v["count"],
            "average": round(v["avg"], 4),
            "min": round(v["min"], 4),
            "max": round(v["max"], 4)
        } for k, v in stats.items()},
        "total_tests": sum(s["count"] for s in stats.values())
    })


@app.get("/api/ocr/accuracy-log")
async def ocr_accuracy_log(engine: str = None, limit: int = 50):
    """Get accuracy log entries."""
    multi_ocr = get_multi_ocr_engine()
    log = multi_ocr.get_accuracy_log(engine, limit)
    
    return JSONResponse({
        "engine": engine,
        "entries": log[-limit:],
        "total": len(log)
    })


@app.post("/evaluate-objectives-run")
def evaluate_objectives_run():
    """Run the objective evaluation suite and persist the latest best selections."""
    try:
        from evaluate_objectives import build_report

        report = build_report()
        save_objective_report(report)
        save_objective_selections(report.get("selections", {}))
        return JSONResponse({
            "message": "Evaluation complete",
            "generated_at": report.get("generated_at"),
            "report_url": "/generated/objective_evaluation_report.json",
            "selection_url": "/generated/objective_selections.json",
        })
    except Exception as exc:
        raise HTTPException(status_code=500, detail=f"Evaluation failed: {exc}") from exc


PROJECT_ROOT = Path(__file__).resolve().parent
TEMPLATES_DIR = PROJECT_ROOT / "templates"
GENERATED_DIR = PROJECT_ROOT / "generated"
STATIC_DIR = PROJECT_ROOT / "static"
GENERATED_DIR.mkdir(parents=True, exist_ok=True)
STATIC_DIR.mkdir(parents=True, exist_ok=True)
# Serve static and generated assets
app.mount("/static", StaticFiles(directory=STATIC_DIR), name="static")
app.mount("/generated", StaticFiles(directory=GENERATED_DIR), name="generated")
SESSION_STORE_PATH = GENERATED_DIR / "chat_sessions.json"

SESSIONS: Dict[str, SessionState] = {}

NDA_MARKERS = {
    "[Name of 1st Party]",
    "[Name of 2nd Party]",
    "[describe the purpose]",
    "[insert]",
}

SET_SINGLE_RE = re.compile(r"^set\s+(\[[^\]]+\])\s*=\s*(.+)$", re.IGNORECASE)
SET_INDEXED_RE = re.compile(r"^set\s+(\[[^\]]+\])#(\d+)\s*=\s*(.+)$", re.IGNORECASE)


def _scan_top_level_placeholders(text: str) -> List[Dict[str, Any]]:
    """Return top-level [...] placeholders with start/end offsets, supporting nested brackets."""
    results: List[Dict[str, Any]] = []
    source = text or ""
    depth = 0
    start: Optional[int] = None

    for i, ch in enumerate(source):
        if ch == "[":
            if depth == 0:
                start = i
            depth += 1
            continue

        if ch == "]" and depth > 0:
            depth -= 1
            if depth == 0 and start is not None:
                results.append({"token": source[start : i + 1], "start": start, "end": i + 1})
                start = None

    return results


def _extract_occurrences_in_order(template_path: Path) -> List[Dict[str, str]]:
    doc = Document(str(template_path))
    rows: List[Dict[str, str]] = []

    def scan(text: str) -> None:
        line = text or ""
        for item in _scan_top_level_placeholders(line):
            start, end = int(item["start"]), int(item["end"])
            rows.append(
                {
                    "token": str(item["token"]),
                    "before": _clean_context_text(line[max(0, start - 80) : start]),
                    "after": _clean_context_text(line[end : min(len(line), end + 80)]),
                    "line": _clean_context_text(line),
                }
            )

    for paragraph in doc.paragraphs:
        scan(paragraph.text)

    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    scan(paragraph.text)

    return rows


def extract_placeholders_with_counts_balanced(template_path: Path) -> Counter:
    counts: Counter = Counter()
    for row in _extract_occurrences_in_order(template_path):
        counts[row["token"]] += 1
    return counts


def _session_to_dict(state: SessionState) -> Dict[str, Any]:
    if hasattr(state, "model_dump"):
        return state.model_dump()
    return state.dict()


def persist_sessions() -> None:
    payload = {session_id: _session_to_dict(state) for session_id, state in SESSIONS.items()}
    SESSION_STORE_PATH.write_text(json.dumps(payload, indent=2), encoding="utf-8")


def restore_sessions() -> None:
    if not SESSION_STORE_PATH.exists():
        return

    raw = json.loads(SESSION_STORE_PATH.read_text(encoding="utf-8"))
    for session_id, state_payload in raw.items():
        SESSIONS[session_id] = SessionState(**state_payload)


def list_templates() -> List[str]:
    return sorted([file.name for file in TEMPLATES_DIR.glob("*.docx")])


def get_session_or_404(session_id: str) -> SessionState:
    state = SESSIONS.get(session_id)
    if not state:
        raise HTTPException(status_code=404, detail="Session not found")
    return state


def _smart_title(value: str) -> str:
    words = [w for w in re.split(r"\s+", value.strip()) if w]
    titled = " ".join(w.capitalize() for w in words)
    return titled.replace("Ceo", "CEO").replace("Vat", "VAT")


def _clean_context_text(text: str) -> str:
    text = re.sub(r"\s+", " ", text or "").strip()
    return text.strip(" -:,.\"")


def _infer_insert_label(before: str, after: str, occurrence: int) -> str:
    before_clause = re.split(r"[,;]", before or "")[-1].strip().lower()
    after_clause = re.split(r"[,;]", after or "")[0].strip().lower()
    context = f"{before_clause} {after_clause}".strip()

    # Use immediate words before placeholder first; this is usually the true field label.
    before_priority = [
        (r"is dated|dated", "Agreement date"),
        (r"registry|registration", "Company registry code"),
        (r"personal identification|identity", "Personal identification code"),
        (r"e-?mail", "Email address"),
        (r"address", "Address"),
        (r"must report to|report to", "Reporting manager name or title"),
        (r"vat|tax", "Tax or VAT number"),
        (r"salary|monthly gross|gross salary|wage", "Gross monthly salary (EUR)"),
        (r"commence|start date|employment starts", "Employment start date"),
        (r"probation", "Probation period value"),
        (r"annual leave|holiday", "Annual leave days"),
        (r"working hours|hours per week", "Working hours"),
        (r"job title|position", "Job title"),
        (r"place of work|workplace", "Place of work"),
        (r"no less than", "Notice period value"),
        (r"at least", "Minimum period or amount"),
        (r"eur|fee|per hour|per day|remuneration", "Fee amount"),
        (r"within", "Payment term in days"),
    ]
    for pattern, label in before_priority:
        if re.search(pattern, before_clause):
            return label

    # Then fallback to the immediate neighboring phrase.
    local_priority = [
        (r"is dated|dated", "Agreement date"),
        (r"registry|registration", "Company registry code"),
        (r"personal identification|identity", "Personal identification code"),
        (r"e-?mail", "Email address"),
        (r"address", "Address"),
        (r"must report to|report to", "Reporting manager name or title"),
        (r"vat|tax", "Tax or VAT number"),
        (r"salary|monthly gross|gross salary|wage", "Gross monthly salary (EUR)"),
        (r"commence|start date|employment starts", "Employment start date"),
        (r"probation", "Probation period value"),
        (r"annual leave|holiday", "Annual leave days"),
        (r"working hours|hours per week", "Working hours"),
        (r"job title|position", "Job title"),
        (r"place of work|workplace", "Place of work"),
        (r"no less than|written notice|terminate", "Notice period value"),
        (r"at least", "Minimum period or amount"),
        (r"eur|fee|per hour|per day|remuneration", "Fee amount"),
        (r"within", "Payment term in days"),
    ]
    for pattern, label in local_priority:
        if re.search(pattern, context):
            return label

    broader_context = f"{before} {after}".lower()
    keyword_map = [
        (r"registry|registration", "Company registry code"),
        (r"personal identification|identity", "Personal identification code"),
        (r"vat|tax", "Tax or VAT number"),
        (r"e-?mail", "Email address"),
        (r"address", "Address"),
        (r"must report to|report to", "Reporting manager name or title"),
        (r"salary|monthly gross|gross salary|wage", "Gross monthly salary (EUR)"),
        (r"commence|start date|employment starts", "Employment start date"),
        (r"probation", "Probation period value"),
        (r"annual leave|holiday", "Annual leave days"),
        (r"working hours|hours per week", "Working hours"),
        (r"job title|position", "Job title"),
        (r"place of work|workplace", "Place of work"),
        (r"dated|agreement is dated", "Agreement date"),
        (r"no less than|written notice|terminate", "Notice period value"),
        (r"within\s+\[?insert\]?\s+days|within\s+\d+\s+days", "Payment term in days"),
        (r"fee of eur|per hour|per day|remuneration", "Fee amount"),
        (r"description of the services", "Services description"),
    ]
    for pattern, label in keyword_map:
        if re.search(pattern, broader_context):
            return label

    # Final fallback: derive a short phrase from nearby words instead of "Detail N".
    candidate = re.sub(r"[^a-zA-Z0-9\s-]", " ", before_clause).strip()
    words = [w for w in candidate.split() if len(w) > 2]
    if words:
        return _smart_title(" ".join(words[-4:]))
    return "Additional detail"


def _normalize_clause_text(value: str) -> str:
    text = (value or "").strip()
    text = re.sub(r"\[(insert|date|dd\s+month\s+yyyy)(?![^\]]*\])", r"[\1]", text, flags=re.IGNORECASE)
    if text.count("[") > text.count("]"):
        text += "]" * (text.count("[") - text.count("]"))
    return text


def _parse_option_clause_token(token: str) -> Optional[Dict[str, str]]:
    match = re.match(r"^\[\s*OPTION\s+(\d+)\s*:\s*(.*)\]$", token.strip(), re.IGNORECASE | re.DOTALL)
    if not match:
        optional_match = re.match(r"^\[\s*OPTIONAL\s*:\s*(.*)\]$", token.strip(), re.IGNORECASE | re.DOTALL)
        if not optional_match:
            return None
        return {
            "option_no": "optional",
            "clause_text": _normalize_clause_text(optional_match.group(1).strip()),
        }
    return {
        "option_no": match.group(1),
        "clause_text": _normalize_clause_text(match.group(2).strip()),
    }


def _build_label_from_occurrence(token: str, before: str, after: str, occurrence: int) -> str:
    cleaned = token.strip()[1:-1].strip() if token.startswith("[") and token.endswith("]") else token
    low = cleaned.lower()

    if low == "insert":
        return _infer_insert_label(before, after, occurrence)

    if low in {"date", "dd month yyyy"}:
        return "Agreement date"

    if "representative" in low and "title" in low:
        return "Representative title"

    if "representative" in low and "name" in low:
        return "Representative name"

    # Normalize instructional prefixes used in many templates.
    cleaned = re.sub(r"^(insert|add)\s+", "", cleaned, flags=re.IGNORECASE).strip()
    cleaned = re.sub(r"\s+where this is not self-explanatory\.?$", "", cleaned, flags=re.IGNORECASE).strip()

    low_cleaned = cleaned.lower()
    if "job title" in low_cleaned or "position" in low_cleaned:
        return "Job title"
    if "reason" in low_cleaned and "definite term" in low_cleaned:
        return "Reason for definite-term employment"
    if "date or the event" in low_cleaned or ("event" in low_cleaned and "employment" in low_cleaned):
        return "Employment end date or event"

    readable = re.sub(r"[_-]+", " ", cleaned)
    readable = re.sub(r"\s+", " ", readable).strip(" ':.\"")
    return _smart_title(readable) if readable else f"Field {occurrence}"


def _infer_choice_prompt(slash_values: List[str], before: str, after: str) -> str:
    joined = f"{before} {after}".lower()
    normalized_values = [v.lower() for v in slash_values]

    if "written notice" in joined or "no less than" in joined or "terminate" in joined:
        return "Choose the notice period unit"

    if "fee" in joined or "per" in joined or "time units" in joined or "invoice" in joined:
        return "Choose the billing unit"

    if normalized_values == ["days", "weeks", "months"] or set(normalized_values) == {"days", "weeks", "months"}:
        return "Choose the duration unit"

    return "Choose one option"


def _annotate_display_progress(queue: List[Dict[str, Any]]) -> None:
    label_totals: Dict[str, int] = {}
    label_seen: Dict[str, int] = {}
    for q in queue:
        if q.get("kind") not in {"default", "choice"}:
            continue
        key = f"{q.get('kind')}::{q.get('label', '')}"
        label_totals[key] = label_totals.get(key, 0) + 1

    for q in queue:
        if q.get("kind") not in {"default", "choice"}:
            continue
        key = f"{q.get('kind')}::{q.get('label', '')}"
        total = label_totals.get(key, 1)
        if total > 1:
            label_seen[key] = label_seen.get(key, 0) + 1
            q["display_index"] = label_seen[key]
            q["display_total"] = total


def queue_from_template(template_path: Path) -> List[Dict[str, Any]]:
    occurrences = _extract_occurrences_in_order(template_path)
    base_rows: List[Dict[str, str]] = []
    option_rows: List[Dict[str, Any]] = []

    for row in occurrences:
        parsed = _parse_option_clause_token(row["token"])
        if parsed:
            inner = [item["token"] for item in _scan_top_level_placeholders(parsed["clause_text"])]
            option_rows.append(
                {
                    "token": row["token"],
                    "kind": "option_clause",
                    "index": None,
                    "label": (
                        "Optional clause"
                        if parsed["option_no"] == "optional"
                        else f"Clause option {parsed['option_no']}"
                    ),
                    "clause_text": parsed["clause_text"],
                    "clause_preview": parsed["clause_text"][:140] + ("..." if len(parsed["clause_text"]) > 140 else ""),
                    "options": ["Keep clause", "Remove clause"],
                    "inner_tokens": inner,
                }
            )
        else:
            base_rows.append(row)

    total_per_token: Dict[str, int] = {}
    for row in base_rows:
        token = row["token"]
        total_per_token[token] = total_per_token.get(token, 0) + 1

    seen_per_token: Dict[str, int] = {}
    queue: List[Dict[str, Any]] = []
    for row in base_rows:
        token = row["token"]
        seen_per_token[token] = seen_per_token.get(token, 0) + 1
        idx = seen_per_token[token] if total_per_token.get(token, 0) > 1 else None
        queue.append(build_generic_question_item(token, idx, row["before"], row["after"]))

    # Ask clause keep/remove decisions after core fields for cleaner flow.
    queue.extend(option_rows)
    _annotate_display_progress(queue)
    return queue


def build_generic_question_item(token: str, index: Optional[int], before: str = "", after: str = "") -> Dict[str, Any]:
    cleaned = token.strip()[1:-1].strip() if token.startswith("[") and token.endswith("]") else token

    # Ambiguous slash placeholders like "days' / weeks' / months'" should be asked as choices.
    slash_values = [part.strip(" '\"") for part in cleaned.split("/")]
    slash_values = [part for part in slash_values if part]
    if len(slash_values) >= 2 and all(len(part) <= 20 for part in slash_values):
        return {
            "token": token,
            "index": index,
            "kind": "choice",
            "label": "Select one option",
            "prompt": _infer_choice_prompt(slash_values, before, after),
            "options": slash_values,
        }

    label = _build_label_from_occurrence(token, before, after, index or 1)
    readable = label if len(label) <= 72 else (label[:72] + "...")
    return {
        "token": token,
        "index": index,
        "kind": "default",
        "label": readable,
        "options": [],
    }


def queue_for_nda_guided_flow() -> List[Dict[str, Any]]:
    return [
        {"token": "[dd Month YYYY]", "index": None, "prompt": "Agreement date (example: 31 March 2026)", "options": []},
        {"token": "[Name of 1st Party]", "index": None, "prompt": "Name of 1st party", "options": []},
        {"token": "[Name of 2nd Party]", "index": None, "prompt": "Name of 2nd party", "options": []},
        {"token": "[Name of one Party]", "index": None, "prompt": "Short repeat name for one Party (usually same as 1st party)", "options": []},
        {"token": "[Name of the other Party]", "index": None, "prompt": "Short repeat name for the other Party (usually same as 2nd party)", "options": []},
        {"token": "[describe the purpose]", "index": None, "prompt": "Purpose of disclosure", "options": ["evaluating a strategic partnership", "commercial due diligence", "technical integration discussions", "potential investment discussions"]},
        {"token": "[other Party's domicile and type]", "index": None, "prompt": "Other party domicile and legal type", "options": ["a French private limited company with registered office in Paris, France", "an Estonian private limited company with registered office in Tallinn, Estonia", "a UK private company limited by shares"]},
        {"token": "[Representative's name]", "index": 1, "prompt": "Representative name for 1st party", "options": []},
        {"token": "[Representative's title]", "index": 1, "prompt": "Representative title for 1st party", "options": ["CEO", "Managing Director", "Director", "Founder"]},
        {"token": "[Representative's name]", "index": 2, "prompt": "Representative name for 2nd party", "options": []},
        {"token": "[Representative's title]", "index": 2, "prompt": "Representative title for 2nd party", "options": ["CEO", "Managing Director", "Director", "Founder"]},
        {"token": "[insert]", "index": 1, "prompt": "1st party registry number", "options": []},
        {"token": "[insert]", "index": 2, "prompt": "2nd party registry number", "options": []},
        {"token": "[insert]", "index": 3, "prompt": "1st party address", "options": []},
        {"token": "[insert]", "index": 4, "prompt": "2nd party address", "options": []},
        {"token": "[insert]", "index": 5, "prompt": "1st party email", "options": []},
        {"token": "[insert]", "index": 6, "prompt": "2nd party email", "options": []},
        {"token": "[insert]", "index": 7, "prompt": "1st party tax or VAT number", "options": []},
        {"token": "[insert]", "index": 8, "prompt": "2nd party tax or VAT number", "options": []},
    ]


def queue_from_missing_tokens(state: SessionState, tokens: List[str]) -> List[Dict[str, Any]]:
    queue: List[Dict[str, Any]] = []
    seen = set()

    for token in tokens:
        if token in seen:
            continue
        seen.add(token)

        count = state.placeholders.get(token, 1)
        existing = state.data.get(token)

        if isinstance(existing, list):
            missing_indices: List[int] = []
            for idx in range(1, count + 1):
                current = existing[idx - 1] if idx - 1 < len(existing) else ""
                if not str(current).strip() or str(current).strip().upper() == "MISSING":
                    missing_indices.append(idx)
            if not missing_indices:
                missing_indices = list(range(1, count + 1))

            for idx in missing_indices:
                queue.append({"token": token, "index": idx})
            continue

        if count <= 1:
            queue.append({"token": token, "index": None})
            continue

        for idx in range(1, count + 1):
            queue.append({"token": token, "index": idx})

    return queue


def _clean_token_for_display(token: str) -> str:
    """Extract a clean, short label from a raw token."""
    cleaned = token.strip()[1:-1].strip() if token.startswith("[") and token.endswith("]") else token
    low = cleaned.lower()

    semantic_map = [
        ("description of the services", "Services description"),
        ("per hour", "Fee amount"),
        ("per day", "Fee amount"),
        ("fee payable", "Fee amount"),
        ("within", "Payment term in days"),
        ("registry", "Company registry code"),
        ("personal identification", "Personal identification code"),
        ("e-mail", "Email address"),
        ("address", "Address"),
    ]
    for key, label in semantic_map:
        if key in low:
            return label

    # If it's very long, truncate it reasonably
    if len(cleaned) > 60:
        # Try to end at a word boundary
        truncated = cleaned[:57]
        last_space = truncated.rfind(" ")
        if last_space > 20:
            return truncated[:last_space] + "..."
        return truncated + "..."
    return cleaned


def _ordinal(num: int) -> str:
    """Convert a number to a human-friendly ordinal: 1 -> '1st', 2 -> '2nd', etc."""
    if num % 100 in (11, 12, 13):
        suffix = "th"
    elif num % 10 == 1:
        suffix = "st"
    elif num % 10 == 2:
        suffix = "nd"
    elif num % 10 == 3:
        suffix = "rd"
    else:
        suffix = "th"
    return f"{num}{suffix}"


def summarize_preview(data: Dict[str, Any], max_items: int = 8) -> Dict[str, Any]:
    keys = list(data.keys())[:max_items]
    return {key: data[key] for key in keys}


def options_for_question(q: Dict[str, Any], state: SessionState) -> List[str]:
    options = list(q.get("options") or [])
    token = q["token"]
    idx = q["index"]

    if token == "[dd Month YYYY]":
        today = datetime.utcnow().strftime("%d %B %Y")
        options = [today, "31 March 2026"] + [o for o in options if o not in {today, "31 March 2026"}]

    if token == "[Name of one Party]":
        party1 = state.data.get("[Name of 1st Party]")
        if isinstance(party1, str) and party1.strip():
            options = [party1] + [o for o in options if o != party1]

    if token == "[Name of the other Party]":
        party2 = state.data.get("[Name of 2nd Party]")
        if isinstance(party2, str) and party2.strip():
            options = [party2] + [o for o in options if o != party2]

    if token == "[insert]" and idx == 6:
        first_email_list = state.data.get("[insert]")
        if isinstance(first_email_list, list) and len(first_email_list) >= 5:
            first_email = str(first_email_list[4]).strip()
            if "@" in first_email:
                domain = first_email.split("@", 1)[1]
                options = [f"info@{domain}", f"legal@{domain}"] + options

    dedup = []
    seen = set()
    for option in options:
        value = str(option).strip()
        if not value:
            continue
        low = value.lower()
        if low in seen:
            continue
        seen.add(low)
        dedup.append(value)
    return dedup[:6]


def next_missing_question(state: SessionState) -> str:
    if state.pending_missing_index >= len(state.pending_missing_queue):
        state.stage = "ready_to_generate"
        persist_sessions()
        return "Great, all missing fields are fixed. Type 'generate' to create the document."

    q = state.pending_missing_queue[state.pending_missing_index]
    token = q["token"]
    idx = q["index"]
    total = len(state.pending_missing_queue)
    step_text = f"({state.pending_missing_index + 1} of {total})"
    
    if idx is None:
        label = _clean_token_for_display(token)
        return f"Please provide: {label} {step_text}"
    
    label = _clean_token_for_display(token)
    count = state.placeholders.get(token, 1)
    ordinal = _ordinal(idx)
    return f"Please provide: {label} ({ordinal}, {idx}/{count}) {step_text}"


def next_nda_question(state: SessionState) -> ChatMessageResponse:
    if state.question_index >= len(state.question_queue):
        state.stage = "ready_to_generate"
        persist_sessions()
        return ChatMessageResponse(
            session_id=state.session_id,
            stage=state.stage,
            reply="Perfect. All NDA fields are collected. Type 'generate' to create your document.",
            data_preview=summarize_preview(state.data),
            can_generate=True,
        )

    q = state.question_queue[state.question_index]
    prompt = q.get("prompt") or "Please provide a value"
    choices = options_for_question(q, state)
    prompt_text = f"Step {state.question_index + 1}/{len(state.question_queue)}: {prompt}"

    return ChatMessageResponse(
        session_id=state.session_id,
        stage=state.stage,
        reply=prompt_text,
        options=choices if choices else None,
        can_generate=False,
    )


def next_generic_question(state: SessionState) -> str:
    if state.question_index >= len(state.question_queue):
        state.stage = "ready_to_generate"
        return (
            "Nice, I have everything I need. Type 'generate' to create your document."
        )

    q = state.question_queue[state.question_index]
    token = q["token"]
    idx = q["index"]
    step_info = f"[Step {state.question_index + 1}/{len(state.question_queue)}]"

    if q.get("kind") == "option_clause":
        label = q.get("label", "Clause option")
        preview = q.get("clause_preview", "")
        # Count total clause option questions to show progress
        clause_count = sum(1 for qi in state.question_queue if qi.get("kind") == "option_clause")
        clause_number = sum(1 for qi in state.question_queue[:state.question_index] if qi.get("kind") == "option_clause") + 1
        return (
            f"{step_info} Clause decision {clause_number}/{clause_count}.\n"
            f"Keep this clause?\n"
            f"{preview}\n"
            "Use the buttons: Keep clause or Remove clause."
        )

    if q.get("kind") == "choice":
        prompt = q.get("prompt", "Please choose one option")
        display_index = q.get("display_index")
        display_total = q.get("display_total")
        if idx is None:
            return f"{step_info} {prompt}"
        if display_index and display_total:
            ordinal = _ordinal(int(display_index))
            return f"{step_info} {prompt} ({ordinal}, {display_index}/{display_total})"
        count = state.placeholders.get(token, 1)
        ordinal = _ordinal(idx)
        return f"{step_info} {prompt} ({ordinal}, {idx}/{count})"

    label = q.get("label", token)
    display_index = q.get("display_index")
    display_total = q.get("display_total")
    if idx is None:
        return f"{step_info} Please provide: {label}"
    if display_index and display_total:
        ordinal = _ordinal(int(display_index))
        return f"{step_info} Please provide: {label} ({ordinal}, {display_index}/{display_total})"
    return f"{step_info} Please provide: {label}"


def options_for_generic_question(state: SessionState) -> Optional[List[str]]:
    if state.question_index >= len(state.question_queue):
        return None

    q = state.question_queue[state.question_index]
    options = list(q.get("options") or [])

    # Friendly defaults for common fields in generic templates.
    label = str(q.get("label", "")).lower()
    if "date" in label:
        today = datetime.utcnow().strftime("%d %B %Y")
        options = [today] + options
    elif re.search(r"\btitle\b", label):
        options = options + ["CEO", "Managing Director", "Director", "Founder"]

    dedup = []
    seen = set()
    for option in options:
        value = str(option).strip()
        if not value:
            continue
        low = value.lower()
        if low in seen:
            continue
        seen.add(low)
        dedup.append(value)
    return dedup[:6] or None


def normalize_generic_answer(state: SessionState, raw_text: str) -> str:
    q = state.question_queue[state.question_index]
    text = raw_text.strip()

    if q.get("kind") != "option_clause":
        return text

    lower = text.lower()
    clause_text = str(q.get("clause_text", "")).strip()
    if lower in {"keep clause", "keep", "yes", "use"}:
        return _normalize_clause_text(clause_text)
    if lower in {"remove clause", "remove", "no", "skip", "none"}:
        return ""
    return "__INVALID_OPTION_DECISION__"


def _max_index_for_token(queue: List[Dict[str, Any]], token: str) -> int:
    max_idx = 0
    for item in queue:
        if item.get("token") != token:
            continue
        idx = item.get("index")
        if isinstance(idx, int) and idx > max_idx:
            max_idx = idx
    return max_idx


def _upgrade_scalar_token_to_indexed(state: SessionState, token: str) -> None:
    entries = [item for item in state.question_queue if item.get("token") == token]
    if not entries:
        return

    none_entries = [item for item in entries if item.get("index") is None]
    if not none_entries:
        return

    # Convert existing None-index entries to explicit sequential indexes.
    seq = 1
    for item in entries:
        if item.get("index") is None:
            item["index"] = seq
            seq += 1

    current = state.data.get(token)
    if current is None:
        return

    if isinstance(current, list):
        return

    state.data[token] = [str(current)]


def append_clause_followups_for_kept_option(state: SessionState, option_question: Dict[str, Any]) -> None:
    inner_tokens = [str(t) for t in option_question.get("inner_tokens") or []]
    if not inner_tokens:
        return

    clause_text = str(option_question.get("clause_text", ""))
    for token in inner_tokens:
        _upgrade_scalar_token_to_indexed(state, token)
        next_idx = _max_index_for_token(state.question_queue, token) + 1
        item = build_generic_question_item(token, next_idx, clause_text, "")
        state.question_queue.append(item)
        state.placeholders[token] = max(state.placeholders.get(token, 0), next_idx)

    _annotate_display_progress(state.question_queue)


def apply_set_command(state: SessionState, text: str) -> Optional[str]:
    indexed = SET_INDEXED_RE.match(text)
    if indexed:
        token, raw_index, value = indexed.groups()
        if token not in state.data or not isinstance(state.data[token], list):
            return f"Cannot set indexed value. Placeholder {token} is not a list field."

        index = int(raw_index) - 1
        if index < 0 or index >= len(state.data[token]):
            return f"Index out of range for {token}."

        state.data[token][index] = value.strip()
        persist_sessions()
        return f"Updated {token}#{raw_index}."

    single = SET_SINGLE_RE.match(text)
    if single:
        token, value = single.groups()
        existing = state.data.get(token)
        if isinstance(existing, list):
            return (
                f"{token} is a list field. Use: set {token}#N=value "
                "(example: set [insert]#3=Tallinn, Estonia)."
            )

        state.data[token] = value.strip()
        persist_sessions()
        return f"Updated {token}."

    return None


def generate_for_session(state: SessionState) -> GenerateResponse:
    if not state.template_path:
        raise HTTPException(status_code=400, detail="No template selected")

    if not state.data:
        raise HTTPException(status_code=400, detail="No data available to generate document")

    stamp = datetime.utcnow().strftime("%Y%m%d_%H%M%S")
    stem = Path(state.template_name or "document").stem
    out_doc = GENERATED_DIR / f"api_{stem}_{state.session_id[:8]}_{stamp}.docx"

    generator = TemplateDocumentGenerator()
    document_path = generator.generate(
        template_path=state.template_path,
        output_path=str(out_doc),
        data=state.data,
        strict=True,
    )

    report_path = None
    warnings: List[str] = []
    if state.is_nda_like and state.report:
        warnings = state.report.get("warnings", [])
        out_report = GENERATED_DIR / f"api_{stem}_{state.session_id[:8]}_{stamp}_report.json"
        filler = IntelligentNDAFiller(template_path=state.template_path)
        report_for_save = dict(state.report)
        report_for_save["placeholder_data"] = state.data
        filler.save_report(report_for_save, str(out_report))
        report_path = str(out_report)

    state.stage = "completed"
    persist_sessions()
    return GenerateResponse(
        session_id=state.session_id,
        document_path=document_path,
        report_path=report_path,
        warnings=warnings,
    )


@app.get("/api/health")
def health() -> Dict[str, str]:
    return {"status": "ok"}


def _generated_url(path_value: Optional[str]) -> Optional[str]:
    if not path_value:
        return None
    path = Path(path_value).resolve()
    try:
        rel = path.relative_to(GENERATED_DIR.resolve())
    except ValueError:
        return None
    return "/generated/" + rel.as_posix()


def _ocr_with_selected_engine(image_path: Path, save_annotated_to: Optional[Path] = None) -> Dict[str, Any]:
    analyzer = DocumentAnalyzer()
    return analyzer.analyze(str(image_path), save_annotated_to=str(save_annotated_to) if save_annotated_to else None)


@app.post("/api/ocr/analyze", response_model=OCRAnalyzeResponse)
def ocr_analyze(request: OCRAnalyzeRequest) -> OCRAnalyzeResponse:
    raw_path = Path(request.image_path)
    image_path = raw_path if raw_path.is_absolute() else (PROJECT_ROOT / raw_path)
    image_path = image_path.resolve()

    if not image_path.exists():
        raise HTTPException(status_code=404, detail=f"Image not found: {image_path}")

    annotated_path: Optional[Path] = None
    if request.save_annotated:
        stamp = datetime.utcnow().strftime("%Y%m%d_%H%M%S")
        annotated_path = GENERATED_DIR / f"ocr_annotated_{stamp}.png"

    try:
        result = _ocr_with_selected_engine(image_path, save_annotated_to=annotated_path)
    except Exception as exc:
        raise HTTPException(
            status_code=500,
            detail=(
                "OCR analysis failed. Ensure OCR dependencies are installed "
                "(opencv-python, easyocr, spacy, and model en_core_web_sm). "
                f"Error: {exc}"
            ),
        ) from exc

    return OCRAnalyzeResponse(
        image_path=str(image_path),
        text=str(result.get("text", "")),
        entities=list(result.get("entities", [])),
        keyword_hits=list(result.get("keyword_hits", [])),
        ocr_results=list(result.get("ocr_results", [])),
        annotated_image_path=result.get("annotated_image_path"),
        annotated_image_url=_generated_url(result.get("annotated_image_path")),
    )


@app.post("/api/ocr/analyze-upload", response_model=OCRAnalyzeResponse)
async def ocr_analyze_upload(file: UploadFile = File(...), save_annotated: bool = True) -> OCRAnalyzeResponse:
    if not file.filename:
        raise HTTPException(status_code=400, detail="Missing uploaded filename")

    ext = Path(file.filename).suffix.lower()
    if ext not in {".png", ".jpg", ".jpeg", ".webp", ".bmp", ".tif", ".tiff"}:
        raise HTTPException(status_code=400, detail="Unsupported file type. Upload an image file.")

    upload_dir = GENERATED_DIR / "ocr_uploads"
    upload_dir.mkdir(parents=True, exist_ok=True)
    stamp = datetime.utcnow().strftime("%Y%m%d_%H%M%S_%f")
    image_path = upload_dir / f"upload_{stamp}{ext}"
    image_path.write_bytes(await file.read())

    annotated_path: Optional[Path] = None
    if save_annotated:
        annotated_path = GENERATED_DIR / f"ocr_annotated_{stamp}.png"

    try:
        result = _ocr_with_selected_engine(image_path, save_annotated_to=annotated_path)
    except Exception as exc:
        raise HTTPException(
            status_code=500,
            detail=(
                "OCR analysis failed. Ensure OCR dependencies are installed "
                "(opencv-python, easyocr, spacy, and model en_core_web_sm). "
                f"Error: {exc}"
            ),
        ) from exc

    return OCRAnalyzeResponse(
        image_path=str(image_path),
        text=str(result.get("text", "")),
        entities=list(result.get("entities", [])),
        keyword_hits=list(result.get("keyword_hits", [])),
        ocr_results=list(result.get("ocr_results", [])),
        annotated_image_path=result.get("annotated_image_path"),
        annotated_image_url=_generated_url(result.get("annotated_image_path")),
    )


@lru_cache(maxsize=1)
def _load_document_intelligence() -> LegalDocumentIntelligence:
    return LegalDocumentIntelligence()


def _build_document_intelligence_response(payload: Dict[str, Any]) -> DocumentIntelligenceResponse:
    return DocumentIntelligenceResponse(
        source_name=payload.get("source_name"),
        document_kind=str(payload.get("document_kind", "general document")),
        summary=str(payload.get("summary", "")),
        domain_tags=list(payload.get("domain_tags", [])),
        key_clauses=list(payload.get("key_clauses", [])),
        startup_signals=list(payload.get("startup_signals", [])),
        risk_flags=list(payload.get("risk_flags", [])),
        open_questions=list(payload.get("open_questions", [])),
        answer=payload.get("answer"),
        evidence=list(payload.get("evidence", [])),
    )


@app.post("/api/document/intel", response_model=DocumentIntelligenceResponse)
def document_intel(request: DocumentIntelligenceRequest) -> DocumentIntelligenceResponse:
    analyzer = _load_document_intelligence()

    try:
        if request.document_path and request.document_path.strip():
            payload = analyzer.analyze_text(
                analyzer.load_text_from_path(request.document_path),
                source_name=request.source_name or Path(request.document_path).name,
                question=request.question,
            )
            return _build_document_intelligence_response(payload)

        if request.text and request.text.strip():
            payload = analyzer.analyze_text(
                request.text,
                source_name=request.source_name or "text input",
                question=request.question,
            )
            return _build_document_intelligence_response(payload)
    except RuntimeError as error:
        raise HTTPException(status_code=503, detail=str(error)) from error

    raise HTTPException(status_code=400, detail="Provide either text or document_path")


@app.post("/api/document/intel-upload", response_model=DocumentIntelligenceResponse)
async def document_intel_upload(file: UploadFile = File(...), question: str = "") -> DocumentIntelligenceResponse:
    if not file.filename:
        raise HTTPException(status_code=400, detail="Missing uploaded filename")

    ext = Path(file.filename).suffix.lower()
    image_exts = {".png", ".jpg", ".jpeg", ".webp", ".bmp", ".tif", ".tiff"}
    text_exts = {".docx", ".txt", ".md"}
    if ext not in image_exts and ext not in text_exts:
        raise HTTPException(
            status_code=400,
            detail="Unsupported file type. Upload a document or image file.",
        )

    upload_dir = GENERATED_DIR / "document_intel_uploads"
    upload_dir.mkdir(parents=True, exist_ok=True)
    stamp = datetime.utcnow().strftime("%Y%m%d_%H%M%S_%f")
    document_path = upload_dir / f"document_{stamp}{ext}"
    document_path.write_bytes(await file.read())

    analyzer = _load_document_intelligence()

    text = analyzer.load_text_from_path(str(document_path))

    try:
        payload = analyzer.analyze_text(
            text,
            source_name=file.filename,
            question=question.strip() or None,
        )
        return _build_document_intelligence_response(payload)
    except RuntimeError as error:
        raise HTTPException(status_code=503, detail=str(error)) from error


@app.get("/")
def chat_ui() -> FileResponse:
    ui_path = STATIC_DIR / "home.html"
    if not ui_path.exists():
        ui_path = STATIC_DIR / "index.html"
    if not ui_path.exists():
        raise HTTPException(status_code=404, detail="UI file not found")
    return FileResponse(ui_path)


@app.get("/evaluation")
def evaluation_ui() -> FileResponse:
    ui_path = STATIC_DIR / "evaluation.html"
    if not ui_path.exists():
        raise HTTPException(status_code=404, detail="Evaluation UI file not found")
    return FileResponse(ui_path)


@app.get("/api/chat/session/{session_id}")
def get_session(session_id: str) -> Dict[str, Any]:
    state = get_session_or_404(session_id)
    return {
        "session_id": state.session_id,
        "stage": state.stage,
        "template_name": state.template_name,
        "is_nda_like": state.is_nda_like,
        "can_generate": state.stage in {"ready_to_generate"},
        "data_preview": summarize_preview(state.data, max_items=25),
    }


@app.post("/api/chat/start", response_model=StartResponse)
def start_chat() -> StartResponse:
    templates = list_templates()
    if not templates:
        raise HTTPException(status_code=400, detail="No templates found")

    session_id = uuid.uuid4().hex
    SESSIONS[session_id] = SessionState(session_id=session_id, stage="awaiting_template")
    persist_sessions()

    return StartResponse(
        session_id=session_id,
        message="Hi! Pick the document you want, and I will guide you step by step.",
        templates=templates,
    )


@app.post("/api/chat/message", response_model=ChatMessageResponse)
def chat_message(request: ChatMessageRequest) -> ChatMessageResponse:
    state = get_session_or_404(request.session_id)
    text = request.message.strip()

    if state.stage == "awaiting_template":
        templates = list_templates()
        selected = next((name for name in templates if name.lower() == text.lower()), None)
        if not selected:
            return ChatMessageResponse(
                session_id=state.session_id,
                stage=state.stage,
                reply=(
                    "I couldn't find that template. Please click one from the list on the left."
                ),
                can_generate=False,
            )

        template_path = TEMPLATES_DIR / selected
        placeholders = extract_placeholders_with_counts_balanced(template_path)

        if not placeholders:
            raise HTTPException(status_code=400, detail="Selected template has no bracket placeholders")

        state.template_name = selected
        state.template_path = str(template_path)
        state.placeholders = {key: int(value) for key, value in placeholders.items()}
        state.is_nda_like = NDA_MARKERS.issubset(set(placeholders.keys()))

        if state.is_nda_like:
            state.stage = "awaiting_nda_fields"
            state.question_queue = queue_for_nda_guided_flow()
            state.question_index = 0
            state.data = {
                "[Representative's name]": ["", ""],
                "[Representative's title]": ["", ""],
                "[insert]": ["", "", "", "", "", "", "", ""],
            }
            state.report = {"warnings": []}
            persist_sessions()
            return next_nda_question(state)

        state.question_queue = queue_from_template(template_path)
        # Generic mode counts are based on the actual asked queue.
        generic_counts: Dict[str, int] = {}
        for item in state.question_queue:
            token = str(item.get("token", ""))
            idx = item.get("index")
            if not token:
                continue
            if isinstance(idx, int):
                generic_counts[token] = max(generic_counts.get(token, 0), idx)
            else:
                generic_counts[token] = max(generic_counts.get(token, 0), 1)
        state.placeholders = generic_counts
        state.stage = "awaiting_generic_fields"
        persist_sessions()
        return ChatMessageResponse(
            session_id=state.session_id,
            stage=state.stage,
            reply=next_generic_question(state),
            options=options_for_generic_question(state),
            can_generate=False,
        )

    if state.stage == "awaiting_nda_fields":
        q = state.question_queue[state.question_index]
        token = q["token"]
        idx = q["index"]

        if idx is None:
            state.data[token] = text
        else:
            if token not in state.data or not isinstance(state.data[token], list):
                required_len = state.placeholders.get(token, idx)
                state.data[token] = [""] * required_len
            while len(state.data[token]) < idx:
                state.data[token].append("")
            state.data[token][idx - 1] = text

        state.question_index += 1
        persist_sessions()
        return next_nda_question(state)

    if state.stage == "awaiting_generic_fields":
        q = state.question_queue[state.question_index]
        token = q["token"]
        idx = q["index"]
        normalized_text = normalize_generic_answer(state, text)

        if q.get("kind") == "option_clause" and normalized_text == "__INVALID_OPTION_DECISION__":
            return ChatMessageResponse(
                session_id=state.session_id,
                stage=state.stage,
                reply="Please choose one of the two options: Keep clause or Remove clause.",
                options=["Keep clause", "Remove clause"],
                can_generate=False,
            )

        if idx is None:
            existing = state.data.get(token)
            if isinstance(existing, list):
                if not existing:
                    existing.append(normalized_text)
                else:
                    existing[0] = normalized_text
            else:
                state.data[token] = normalized_text
        else:
            if token not in state.data or not isinstance(state.data[token], list):
                existing = state.data.get(token)
                state.data[token] = [] if existing is None else [str(existing)]
            while len(state.data[token]) < idx:
                state.data[token].append("")
            state.data[token][idx - 1] = normalized_text

        if q.get("kind") == "option_clause" and normalized_text:
            append_clause_followups_for_kept_option(state, q)

        state.question_index += 1
        reply = next_generic_question(state)
        persist_sessions()

        return ChatMessageResponse(
            session_id=state.session_id,
            stage=state.stage,
            reply=reply,
            options=options_for_generic_question(state),
            can_generate=state.stage in {"ready_to_generate"},
        )

    if state.stage == "awaiting_missing_fields":
        if state.pending_missing_index >= len(state.pending_missing_queue):
            state.stage = "ready_to_generate"
            persist_sessions()
            return ChatMessageResponse(
                session_id=state.session_id,
                stage=state.stage,
                reply="All missing fields are now filled. Type 'generate'.",
                can_generate=True,
            )

        q = state.pending_missing_queue[state.pending_missing_index]
        token = q["token"]
        idx = q["index"]

        if idx is None:
            state.data[token] = text
        else:
            if token not in state.data or not isinstance(state.data[token], list):
                state.data[token] = [""] * state.placeholders.get(token, idx)
            while len(state.data[token]) < idx:
                state.data[token].append("")
            state.data[token][idx - 1] = text

        state.pending_missing_index += 1
        persist_sessions()

        options: Optional[List[str]] = None
        if state.pending_missing_index < len(state.pending_missing_queue):
            next_q = state.pending_missing_queue[state.pending_missing_index]
            candidate_options = options_for_question(next_q, state)
            options = candidate_options if candidate_options else None

        return ChatMessageResponse(
            session_id=state.session_id,
            stage=state.stage,
            reply=next_missing_question(state),
            options=options,
            can_generate=state.stage in {"ready_to_generate"},
        )

    if state.stage in {"ready_to_generate"}:
        low = text.lower()

        if low == "show":
            return ChatMessageResponse(
                session_id=state.session_id,
                stage=state.stage,
                reply="Current data snapshot.",
                data_preview=state.data,
                can_generate=True,
            )

        update_message = apply_set_command(state, text)
        if update_message:
            persist_sessions()
            return ChatMessageResponse(
                session_id=state.session_id,
                stage=state.stage,
                reply=update_message,
                can_generate=True,
            )

        if low == "generate":
            try:
                result = generate_for_session(state)
            except ValueError as exc:
                message = str(exc)
                if "Unreplaced placeholders:" in message:
                    unresolved_part = message.split("Unreplaced placeholders:", 1)[1]
                    unresolved_part = unresolved_part.split("|", 1)[0].strip()
                    unresolved = [item.strip() for item in unresolved_part.split(",") if item.strip()]

                    state.pending_missing_queue = queue_from_missing_tokens(state, unresolved)
                    state.pending_missing_index = 0
                    state.stage = "awaiting_missing_fields"
                    persist_sessions()

                    return ChatMessageResponse(
                        session_id=state.session_id,
                        stage=state.stage,
                        reply=(
                            "Some placeholders are still empty, so brackets would remain in the document. "
                            + next_missing_question(state)
                        ),
                        can_generate=False,
                    )
                raise

            return ChatMessageResponse(
                session_id=state.session_id,
                stage=state.stage,
                reply=(
                    f"Done. Your document is ready: {result.document_path}. "
                    f"Report: {result.report_path or 'n/a'}."
                ),
                can_generate=False,
            )

        return ChatMessageResponse(
            session_id=state.session_id,
            stage=state.stage,
            reply=(
                "I didn't understand that. Try: 'show', 'set [field]=value', or 'generate'."
            ),
            can_generate=True,
        )

    if state.stage == "completed":
        return ChatMessageResponse(
            session_id=state.session_id,
            stage=state.stage,
            reply="Session already completed. Start a new session with POST /api/chat/start.",
            can_generate=False,
        )

    raise HTTPException(status_code=400, detail=f"Unhandled session stage: {state.stage}")


@app.post("/api/chat/generate", response_model=GenerateResponse)
def generate_document(request: GenerateRequest) -> GenerateResponse:
    state = get_session_or_404(request.session_id)
    if state.stage not in {"ready_to_generate"}:
        raise HTTPException(
            status_code=400,
            detail="Session is not ready for generation. Complete intake first.",
        )
    return generate_for_session(state)


restore_sessions()
app.mount("/static", StaticFiles(directory=str(STATIC_DIR)), name="static")
app.mount("/generated", StaticFiles(directory=str(GENERATED_DIR)), name="generated")


# Endpoint to list available templates
@app.get("/api/templates")
def api_list_templates():
    return list_templates()


if __name__ == "__main__":
    try:
        import uvicorn
    except Exception as exc:
        raise RuntimeError(
            "uvicorn is required to run chatbot_api.py directly. Install project dependencies first."
        ) from exc

    uvicorn.run(app, host="127.0.0.1", port=8000, reload=False)
